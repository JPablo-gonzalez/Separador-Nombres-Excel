import io
import math
import os
import posixpath
import re
import shutil
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass

import streamlit as st
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

# =============================================================================
# 1. DICCIONARIOS SEMÁNTICOS Y CONECTORES
# =============================================================================


def _norm(texto):
    """Mayúsculas y sin tildes/diacríticos: MEJÍA -> MEJIA, MUÑOZ -> MUNOZ."""
    t = unicodedata.normalize("NFD", str(texto).upper())
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


# Conectores que se pegan a la(s) palabra(s) siguiente(s): "DE LA CRUZ", "DEL RIO", "SAN JUAN"...
CONECTORES = {"DE", "DEL", "LA", "LOS", "LAS", "SAN", "SANTA", "VON", "VAN", "DA", "DOS", "DAS", "DI"}
# Los que, al abrir un bloque, delatan un apellido ("De La Cruz", "Del Rio", "Von Trapp").
# SAN y SANTA quedan fuera: también aparecen en nombres.
CONECTORES_APELLIDO = CONECTORES - {"SAN", "SANTA"}

# Bloques religiosos que cierran un nombre compuesto (cuentan como UNA sola palabra y son nombre).
RELIGIOSOS = {
    "DE JESUS", "DEL CARMEN", "DE MARIA", "DEL ROSARIO",
    "DEL PILAR", "DE LOS ANGELES", "DE LOS DOLORES",
}

_NOMBRES_BASE = {
    "JUAN", "CARLOS", "LUIS", "JOSE", "MARIA", "ANDRES", "DAVID", "ALEJANDRO", "ARTURO", "FREDY",
    "EDWIN", "GERARDO", "WILSON", "EDISON", "JHON", "JORGE", "ALEXANDER", "JULIAN", "DIEGO", "DANIEL",
    "MIGUEL", "ANGEL", "HERNANDO", "GUILLERMO", "GUSTAVO", "JAIME", "ALBERTO", "HECTOR", "JAIRO",
    "CESAR", "JULIO", "TITO", "ERNESTO", "NELSON", "CONRADO", "ALIRIO", "LIBANIEL", "EDER", "DIANA",
    "LUCIA", "WILLIAM", "ANTONIO", "ARMANDO", "MIRIAM", "GABRIELA", "ALFONSO", "JOSELIN", "HERIBERTO",
    "REGINA", "AMPARO", "MANUEL", "HERNAN",
    "PABLO", "PEDRO", "SANTIAGO", "SEBASTIAN", "NICOLAS", "FELIPE", "MATEO", "SAMUEL", "JESUS",
    "FERNANDO", "RICARDO", "ROBERTO", "OSCAR", "MAURICIO", "FABIAN", "CRISTIAN", "STEVEN", "BRAYAN",
    "KEVIN", "JHONATAN", "YEISON", "ESTEBAN", "EDUARDO", "ALEXIS", "ANDREA", "PAOLA", "CATALINA",
    "VALENTINA", "SOFIA", "ISABELLA", "CAMILA", "LAURA", "ANA", "CLAUDIA", "PATRICIA", "MARTHA",
    "SANDRA", "LUZ", "ELENA", "YESICA", "LEIDY", "NATALIA", "DANIELA", "MARIANA", "JULIANA",
    "PAULA", "ROSA", "LILIANA", "GLORIA", "BEATRIZ", "ADRIANA", "CAROLINA", "JOHAN", "SERGIO",
    "ISAAC", "EMMANUEL", "THOMAS", "SIMON", "JERONIMO", "MARTIN", "VICTOR", "RAFAEL", "ANDERSON",
    # ampliación
    "CAMILO", "GABRIEL", "LEONARDO", "LEON", "JAVIER", "IVAN", "FELIX", "MARCO", "MARCOS", "MARIO",
    "HUGO", "ORLANDO", "RAUL", "ROBINSON", "ROLANDO", "ELKIN", "EDGAR", "FABIO", "FRANCISCO", "HENRY",
    "JHONNY", "JOAQUIN", "LEONEL", "NESTOR", "OMAR", "OSWALDO", "RODRIGO", "SAUL", "TOMAS", "WALTER",
    "ALVARO", "ALAN", "DYLAN", "JUSTIN", "LUCAS", "BRYAN", "JEFFERSON", "SNEIDER", "YEFERSON", "JHOAN",
    "CRISTHIAN", "STIVEN", "DUVAN", "DEIBY", "YEISSON", "WILMAR", "WILDER", "FAVIO", "FABRICIO",
    "SARA", "NICOLE", "SHIRLEY", "KAREN", "LINA", "MONICA", "JENNIFER", "YULIANA", "MELISSA",
    "ALEJANDRA", "JOHANNA", "JOHANA", "TATIANA", "MARCELA", "ANGELA", "ANGIE", "JUANA", "LUISA",
    "FERNANDA", "ISABEL", "VERONICA", "DAYANA", "ESTEFANIA", "STEFANIA", "MARISOL", "YULIETH",
    "MAYERLY", "ALEXANDRA", "AURA", "BLANCA", "CARMEN", "CECILIA", "CLARA", "DORA", "ELSA", "EMILIA",
    "ESPERANZA", "EVA", "FLOR", "GISELA", "GRACIELA", "INES", "IRENE", "JUDITH", "KATHERINE",
    "LORENA", "LUCERO", "MABEL", "MARGARITA", "MARISELA", "MELANIE", "MERCEDES", "NANCY", "NORA",
    "OLGA", "PILAR", "RAQUEL", "RUTH", "SILVIA", "SOL", "SUSANA", "TERESA", "VANESSA", "VIVIANA",
    "XIMENA", "YOLANDA", "YURANI", "ZULEIDY", "MARIA", "KATHERIN", "KAROL", "LEIDY", "MAIRA",
    "ESTEFANY", "YENNY", "RAMIRO", "CIRO", "YESENIA", "JESSICA", "JESICA", "DAISY", "DEISY", "ERIKA", "ERICA", "SHARON",
}

_APELLIDOS_BASE = {
    "GOMEZ", "ZAPATA", "PEREZ", "OSORIO", "VERA", "BETANCUR", "MORALES", "GALEANO", "ESPINOSA",
    "GUARIN", "CELIS", "RAMIREZ", "HERNANDEZ", "TORO", "OCAMPO", "ARROYAVE", "ARANGO", "MUNOZ",
    "LONDONO", "AGUIRRE", "AMESQUITA", "MARIN", "BETANCURT", "TOBON", "GARCIA", "MEJIA", "ARANZAZU",
    "OSPINA", "SANCHEZ", "GAVIRIA", "CANO", "RUIZ", "BARRERA", "GALLO", "RAMOS", "GRAJALES",
    "GRISALES", "BOTERO", "CASTRO", "BARRETO", "ZAMBRANO", "BUITRAGO", "OBANDO", "GALLEGO", "MESA",
    "ARIAS", "CASTANO", "HERRERA", "MOLINA", "CARDONA", "PARRA", "TASCON", "DIAZ", "LOPEZ", "MACIAS",
    "RODRIGUEZ", "MARTINEZ", "GONZALEZ", "VARGAS", "JIMENEZ", "RESTREPO", "QUINTERO", "FIGUEROA",
    "SUAREZ", "CARDENAS", "ROJAS", "GUTIERREZ", "ORTIZ", "VELEZ", "DUQUE", "VALENCIA", "CORREA",
    "URIBE", "ALZATE", "MORENO", "SALAZAR", "CARDOSO", "MONTOYA", "JARAMILLO", "NARANJO", "PALACIO",
    "PALACIOS", "FRANCO", "HOYOS", "VILLA", "AGUDELO", "VASQUEZ", "TORRES", "FLOREZ", "ESCOBAR",
    "CASTILLO", "ROMERO", "ALVAREZ", "ACEVEDO", "SERNA", "USUGA", "ECHEVERRI", "POSADA", "LOAIZA",
    "HENAO", "MURILLO", "ZULUAGA", "OROZCO", "CARVAJAL", "BEDOYA", "RIVERA", "SIERRA", "PINEDA",
    # ampliación
    "ALVARADO", "ARBOLEDA", "ATEHORTUA", "AVENDANO", "BALLESTEROS", "BERRIO", "BLANDON", "CALLE",
    "CAMPUZANO", "CARO", "CASAS", "CEBALLOS", "CHAVARRIA", "COLORADO", "CORTES", "CUARTAS",
    "DELGADO", "DURANGO", "ESTRADA", "FERNANDEZ", "FUENTES", "GIL", "GIRALDO", "GUERRA", "GUZMAN",
    "HIGUITA", "HINCAPIE", "HOLGUIN", "LEON", "LOZANO", "MADRID", "MARQUEZ", "MAZO", "MEDINA",
    "MENDEZ", "MENDOZA", "MIRA", "MONSALVE", "MONTES", "MORA", "MUNERA", "NIETO", "NUNEZ", "OCHOA",
    "OLARTE", "OSSA", "PABON", "PAEZ", "PATINO", "PELAEZ", "PEREA", "PIEDRAHITA", "QUICENO", "QUIROZ",
    "RENDON", "RENTERIA", "REYES", "RIOS", "RIVAS", "ROLDAN", "SALDARRIAGA", "SOTO", "TABARES",
    "TAMAYO", "TREJOS", "TRUJILLO", "VALDERRAMA", "VALLEJO", "VARELA", "VELASQUEZ", "VERGARA",
    "VIDAL", "YEPES", "CRUZ", "MARTIN", "BEJARANO", "BERMUDEZ", "BUSTAMANTE", "CAICEDO", "CAMACHO",
    "CARRILLO", "CHACON", "CORDOBA", "COSSIO", "DOMINGUEZ", "ESPITIA", "GAITAN", "GALVIS", "GARZON",
    "GIRON", "GUEVARA", "IBARRA", "LARA", "LEAL", "LUNA", "MALDONADO", "MANRIQUE", "MARIN", "MEJIAS",
    "MONTAÑO", "MONTANO", "NAVARRO", "ORDONEZ", "PADILLA", "PAREDES", "PENA", "PRADA", "QUINTANA",
    "RANGEL", "SALGADO", "SANDOVAL", "SEPULVEDA", "SOLANO", "TELLO", "URREGO", "VACA",
    "VEGA", "VILLEGAS", "YEPEZ", "ZABALA", "ZUNIGA",
}

# Los sets exponen el contenido normalizado (sin tildes) para comparar sin importar el acento.
NOMBRES_COMUNES = {_norm(x) for x in _NOMBRES_BASE}
APELLIDOS_COMUNES = {_norm(x) for x in _APELLIDOS_BASE}

ORIENTACIONES = ("N->A", "A->N")
RE_PREFIJO = re.compile(r"^\(\s*([AaBb])\s*\)\s*")


def agrupar_conectores(tokens):
    """Une conectores hispanos con la palabra que les sigue ("DE LA" + "CRUZ" -> "DE LA CRUZ")
    y deja los bloques religiosos finales ("DE JESUS", "DEL CARMEN") como una sola palabra."""
    resultado = []
    n = len(tokens)
    i = 0
    while i < n:
        j = i
        while j < n - 1 and _norm(tokens[j]) in CONECTORES:
            j += 1
        resultado.append(" ".join(tokens[i:j + 1]))
        i = j + 1
    return resultado


# =============================================================================
# 2. MOTOR DE DECISIÓN (orientación y reparto deducidos celda por celda)
# =============================================================================
#
# Para cada celda se enumeran TODAS las lecturas posibles: cada orientación (N->A / A->N) y cada
# reparto de palabras entre nombres (k) y apellidos (m). Cada lectura recibe un puntaje y gana la mayor.
#
#   · Prefijo (A)/(B): fija la orientación de esa celda (y se borra). Solo se decide el reparto.
#   · Celda sin prefijo: las dos orientaciones compiten con la evidencia de las palabras:
#       - diccionarios NOMBRES_COMUNES / APELLIDOS_COMUNES (evidencia fuerte, ±2),
#       - palabras que el propio archivo enseña con sus filas más claras (evidencia media, ±1 a ±1.5),
#       - terminaciones típicas de apellido (-EZ, -IZ, -OZ, -AZ) y conectores iniciales (evidencia débil),
#       - estructura más frecuente (1-2 nombres, 2 apellidos) como último respaldo.
#   · Si una celda no ofrece ninguna evidencia, se usa la tendencia de las filas claras del mismo
#     archivo; si tampoco hay tendencia, el orden natural (Nombres -> Apellidos). Estas filas se
#     marcan para revisión.

MARGEN_ORIENTACION = 1.0      # diferencia mínima entre orientaciones para decidir por evidencia
MARGEN_APRENDER_PREFIJO = 1.5  # con prefijo, el reparto debe ser claro para enseñar palabras nuevas
MARGEN_APRENDER_AUTO = 3.0    # sin prefijo, la lectura debe ser muy clara para enseñar palabras nuevas
MARGEN_REVISAR = 1.0          # por debajo de esta confianza la fila se marca para revisión
ORIENTACION_NATURAL = "N->A"
MAX_ITERACIONES = 4
SUFIJOS_APELLIDO = ("EZ", "IZ", "OZ", "AZ")


@dataclass
class Resultado:
    texto: str                 # texto original de la celda
    p_ape: str = ""
    s_ape: str = ""
    p_nom: str = ""
    s_nom: str = ""
    orientacion: str = ORIENTACION_NATURAL
    origen: str = "AUTO"       # "A", "B" (prefijo) o "AUTO" (deducido)
    confianza: float = 0.0
    revisar: bool = False
    motivo: str = ""


class Lexico:
    """Palabras que el propio archivo enseña, además de los diccionarios base."""

    def __init__(self):
        self.nombres = Counter()
        self.apellidos = Counter()

    def agregar(self, contribucion):
        for (palabra, rol), cuantas in contribucion.items():
            (self.nombres if rol == "N" else self.apellidos)[palabra] += cuantas

    def peso(self, palabra, rol, excluir=None):
        base = NOMBRES_COMUNES if rol == "N" else APELLIDOS_COMUNES
        if palabra in base:
            return 2.0
        veces = (self.nombres if rol == "N" else self.apellidos)[palabra]
        if excluir:                                   # una fila no se enseña a sí misma
            veces -= excluir.get((palabra, rol), 0)
        return min(1.5, 0.75 + 0.25 * veces) if veces > 0 else 0.0


def _afinidades(token, lex, excluir=None):
    """(afinidad como nombre, afinidad como apellido) de una palabra o bloque."""
    clave = _norm(token)
    if clave in RELIGIOSOS:
        return 2.0, -2.0
    palabras = clave.split()
    nucleo = palabras[-1]
    wn = lex.peso(nucleo, "N", excluir)
    wa = lex.peso(nucleo, "A", excluir)
    if wn and wa:                                      # palabra ambigua (León, Martín...)
        an, aa = 0.5, 0.5
    elif wn:
        an, aa = wn, -wn
    elif wa:
        an, aa = -wa, wa
    elif len(nucleo) >= 5 and nucleo.endswith(SUFIJOS_APELLIDO):
        an, aa = -0.75, 0.75
    else:
        an, aa = 0.0, 0.0
    if len(palabras) > 1 and palabras[0] in CONECTORES_APELLIDO:
        an, aa = an - 1.0, aa + 1.0
    return an, aa


def _hipotesis(n):
    """Todas las lecturas posibles: (orientación, nº de nombres k, nº de apellidos m)."""
    if n == 1:
        return [("N->A", 1, 0), ("A->N", 0, 1)]
    return [(o, k, n - k) for o in ORIENTACIONES for k in range(1, n)]


def _prior_estructura(k, m):
    """Estructura más frecuente en Colombia: 2 apellidos y 1-2 nombres."""
    if k + m == 1:
        return 0.0
    return -1.0 * abs(m - 2) - 0.5 * abs(k - 2)


def _indices(n, orientacion, k, m):
    """Posiciones de las palabras que son nombres y de las que son apellidos."""
    if orientacion == "N->A":
        return list(range(k)), list(range(k, n))
    return list(range(m, n)), list(range(m))


def _puntuar(afin, orientacion, k, m):
    ni, ai = _indices(len(afin), orientacion, k, m)
    return (sum(afin[i][0] for i in ni) + sum(afin[i][1] for i in ai) + _prior_estructura(k, m))


@dataclass
class _Evaluacion:
    orientacion: str
    k: int
    m: int
    fallback: bool
    margen_est: float
    margen_or: float

    @property
    def confianza(self):
        return min(self.margen_est, self.margen_or)


def _evaluar(tokens, fija, lex, excluir, prior_archivo):
    n = len(tokens)
    afin = [_afinidades(t, lex, excluir) for t in tokens]
    puntos = {h: _puntuar(afin, *h) for h in _hipotesis(n)}
    mejor = {o: max(p for h, p in puntos.items() if h[0] == o) for o in ORIENTACIONES}

    if fija:
        orientacion, fallback, margen_or = fija, False, math.inf
    else:
        dif = mejor["N->A"] - mejor["A->N"]
        margen_or = abs(dif)
        if margen_or >= MARGEN_ORIENTACION:
            orientacion, fallback = ("N->A" if dif > 0 else "A->N"), False
        else:                                          # sin evidencia: tendencia del archivo u orden natural
            orientacion, fallback = (prior_archivo or ORIENTACION_NATURAL), True

    cand = sorted(((p, h) for h, p in puntos.items() if h[0] == orientacion),
                  key=lambda x: (-x[0], x[1][1]))
    _, (_, k, m) = cand[0]
    margen_est = cand[0][0] - cand[1][0] if len(cand) > 1 else math.inf
    return _Evaluacion(orientacion, k, m, fallback, margen_est, margen_or)


def _contribucion(tokens, ev):
    """Palabras desconocidas de una lectura clara: lo que esa fila le enseña al resto del archivo."""
    ni, ai = _indices(len(tokens), ev.orientacion, ev.k, ev.m)
    contribucion = Counter()
    for rol, indices in (("N", ni), ("A", ai)):
        for i in indices:
            clave = _norm(tokens[i])
            if clave in RELIGIOSOS:
                continue
            nucleo = clave.split()[-1]
            if len(nucleo) < 3 or nucleo in NOMBRES_COMUNES or nucleo in APELLIDOS_COMUNES:
                continue
            contribucion[(nucleo, rol)] += 1
    return contribucion


def _puede_ensenar(ev, fija):
    if fija:
        return ev.margen_est >= MARGEN_APRENDER_PREFIJO
    return (not ev.fallback) and ev.confianza >= MARGEN_APRENDER_AUTO


def separar_prefijo(texto):
    """(texto sin prefijo, orientación fijada por el prefijo o None, origen 'A' | 'B' | 'AUTO')."""
    texto = str(texto).strip()
    m = RE_PREFIJO.match(texto)
    if m:
        letra = m.group(1).upper()
        return texto[m.end():].strip(), ("A->N" if letra == "A" else "N->A"), letra
    return texto, None, "AUTO"


def _armar(tokens, ev):
    ni, ai = _indices(len(tokens), ev.orientacion, ev.k, ev.m)
    nombres = [tokens[i] for i in ni]
    apellidos = [tokens[i] for i in ai]
    return (apellidos[0] if apellidos else "", " ".join(apellidos[1:]),
            nombres[0] if nombres else "", " ".join(nombres[1:]))


def _motivo(origen, ev, prior_archivo):
    flecha = "Nombres → Apellidos" if ev.orientacion == "N->A" else "Apellidos → Nombres"
    if origen in ("A", "B"):
        texto = f"Prefijo ({origen}): {flecha}"
    elif not ev.fallback:
        texto = f"Deducido por los diccionarios: {flecha}"
    elif prior_archivo:
        texto = f"Sin palabras conocidas; se siguió la tendencia del archivo: {flecha}"
    else:
        texto = f"Sin palabras conocidas ni tendencia en el archivo; se asumió el orden natural: {flecha}"
    if ev.margen_est < MARGEN_REVISAR:
        texto += " (reparto de palabras poco claro)"
    return texto


def resolver_lista(textos):
    """Resuelve TODAS las celdas de una columna con un único flujo.

    Devuelve una lista de `Resultado` (uno por texto). Sirve igual para archivos mixtos (celdas con
    (A), con (B) y sin letra) que para archivos sin ninguna letra.
    """
    preparadas = []
    for texto in textos:
        limpio, fija, origen = separar_prefijo(texto)
        preparadas.append((texto, agrupar_conectores(limpio.split()), fija, origen))

    total = len(preparadas)
    lex, contribs, prior = Lexico(), [None] * total, None
    previas, evals = None, []
    for _ in range(MAX_ITERACIONES):
        evals = [(_evaluar(tk, fija, lex, contribs[i], prior) if tk else None)
                 for i, (_, tk, fija, _) in enumerate(preparadas)]
        if previas is not None and all(
                (a is None and b is None) or (a and b and (a.orientacion, a.k) == (b.orientacion, b.k))
                for a, b in zip(previas, evals)):
            break
        previas = evals
        # Se reconstruye lo aprendido a partir de las filas claras de esta vuelta.
        lex, contribs = Lexico(), [None] * total
        conteo = Counter()
        for i, ((_, tk, fija, _), ev) in enumerate(zip(preparadas, evals)):
            if ev is None:
                continue
            if _puede_ensenar(ev, fija):
                contribs[i] = _contribucion(tk, ev)
                lex.agregar(contribs[i])
            if fija is None and not ev.fallback and ev.confianza >= MARGEN_APRENDER_AUTO:
                conteo[ev.orientacion] += 1
        if conteo["N->A"] != conteo["A->N"]:
            prior = "N->A" if conteo["N->A"] > conteo["A->N"] else "A->N"
        else:
            prior = None

    resultados = []
    for (texto, tokens, fija, origen), ev in zip(preparadas, evals):
        if ev is None:
            resultados.append(Resultado(texto=str(texto), origen=origen, motivo="Celda sin nombre"))
            continue
        pa, sa, pn, sn = _armar(tokens, ev)
        resultados.append(Resultado(
            texto=str(texto), p_ape=pa, s_ape=sa, p_nom=pn, s_nom=sn,
            orientacion=ev.orientacion, origen=origen,
            confianza=ev.confianza if math.isfinite(ev.confianza) else 99.0,
            revisar=ev.fallback or ev.confianza < MARGEN_REVISAR,
            motivo=_motivo(origen, ev, prior),
        ))
    return resultados


def analizar_nombre(texto):
    """Una sola celda, sin el contexto del resto del archivo (útil para pruebas)."""
    return resolver_lista([texto])[0]


# =============================================================================
# 3. REFERENCIAS: desplazamiento de columnas (funciones puras sobre texto)
# =============================================================================
# Al insertar columnas, toda referencia a una columna >= ins se desplaza n columnas a la derecha,
# igual que hace Excel. Cada extremo de un rango se evalúa por separado: un rango que cruza el
# punto de inserción se ensancha.

_RE_CELDA = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_RE_COL = re.compile(r"^(\$?)([A-Za-z]{1,3})$")


def _desplazar_col(letras, ins, n):
    c = column_index_from_string(letras.upper())
    return get_column_letter(c + n) if c >= ins else letras


def _desplazar_operando(op, titulo, ins, n, hoja_propia):
    if "!" in op:
        pref, ref = op.rsplit("!", 1)
        if pref.strip("'").replace("''", "'") != titulo:
            return op                      # apunta a otra hoja
    else:
        pref, ref = None, op
        if not hoja_propia:
            return op                      # sin hoja explícita, pero la fórmula es de otra hoja
    if "[" in ref:
        return op                          # referencias estructuradas de tabla (por nombre)
    partes = ref.split(":")
    if len(partes) > 2:
        return op
    nuevas = []
    for p in partes:
        m = _RE_CELDA.match(p)
        if m:
            nuevas.append(f"{m.group(1)}{_desplazar_col(m.group(2), ins, n)}{m.group(3)}{m.group(4)}")
            continue
        m = _RE_COL.match(p)
        if m and len(partes) == 2:         # columnas completas (A:C); solas podrían ser un nombre
            nuevas.append(f"{m.group(1)}{_desplazar_col(m.group(2), ins, n)}")
            continue
        nuevas.append(p)
    ref_nueva = ":".join(nuevas)
    return f"{pref}!{ref_nueva}" if pref is not None else ref_nueva


def _desplazar_formula(texto, titulo, ins, n, hoja_propia):
    """texto incluye el '=' inicial. Si algo falla, devuelve el texto original."""
    try:
        tok = Tokenizer(texto)
        cambio = False
        for t in tok.items:
            if t.type == Token.OPERAND and t.subtype == Token.RANGE:
                nuevo = _desplazar_operando(t.value, titulo, ins, n, hoja_propia)
                if nuevo != t.value:
                    t.value = nuevo
                    cambio = True
        return tok.render() if cambio else texto
    except Exception:
        return texto


def _desplazar_rangos(texto, titulo, ins, n):
    """Lista de rangos separados por espacios (sqref) de la hoja propia."""
    return " ".join(_desplazar_operando(r, titulo, ins, n, True) for r in str(texto).split())


# =============================================================================
# 4. PROCESAMIENTO DIRECTO DEL .XLSX
# =============================================================================
# El .xlsx es un ZIP de archivos XML. En vez de cargarlo entero con openpyxl (que reconstruye el libro,
# crea una celda en memoria por cada fila declarada y pierde formas, imágenes y partes de Office),
# se edita el XML de la hoja en el lugar:
#
#   · solo se recorren las filas que EXISTEN en el archivo (una celda perdida en la fila 1.048.255
#     ya no cuesta un millón de iteraciones);
#   · imágenes, logos, tema, estilos y cualquier otra parte se copian byte a byte;
#   · solo se reescribe lo que cambia al insertar 4 columnas: celdas, anchos, combinaciones,
#     tablas, filtros, fórmulas, formatos condicionales, validaciones, nombres definidos y
#     la posición de los dibujos.

HEADERS_NUEVOS = ["PRIMER APELLIDO", "SEGUNDO APELLIDO", "PRIMER NOMBRE", "SEGUNDO NOMBRE"]
N_NUEVAS = len(HEADERS_NUEVOS)
ANCHO_NUEVAS = 19

LIMITE_SUBIDA_MB = 30            # tamaño máximo del .xlsx subido
LIMITE_DESCOMPRIMIDO_MB = 200    # tamaño máximo del contenido descomprimido (protege de "zip bombs")
FILA_MAX_ENCABEZADO = 30         # el encabezado se busca por encima de esta fila (como antes)


class ColumnaNoEncontrada(Exception):
    pass


class ArchivoNoSoportado(Exception):
    """Archivo inválido o fuera de los límites de seguridad (el mensaje se muestra al usuario)."""


# ---------------------------------------------------------------- utilidades de texto / XML

_RE_ENT_NUM = re.compile(r"&#(x[0-9A-Fa-f]+|[0-9]+);")
_RE_ESC_X = re.compile(r"_x([0-9A-Fa-f]{4})_")
_RE_ATTR = re.compile(rb"""([\w:.-]+)\s*=\s*(?:"([^"]*)"|'([^']*)')""")
_RE_REF_B = re.compile(rb"^([A-Z]+)(\d+)$")
_RE_FILA = re.compile(rb"<row\b([^>]*?)(?:/>|>(.*?)</row>)", re.S)
_RE_CEL = re.compile(rb"<c\b([^>]*?)(?:/>|>(.*?)</c>)", re.S)
_RE_F = re.compile(rb"<f\b([^>]*?)(?:/>|>(.*?)</f>)", re.S)
_RE_T = re.compile(rb"<t\b[^>]*?(?:/>|>(.*?)</t>)", re.S)
_RE_RPH = re.compile(rb"<rPh\b.*?</rPh>", re.S)
_RE_SPANS = re.compile(rb'\sspans="[^"]*"')
_RE_R_ATTR = re.compile(rb'(\s)r="[^"]*"')
_RE_R_CELDA = re.compile(rb'\sr="([A-Z]+)\d+"')
_RE_S_ATTR = re.compile(rb'\ss="(\d+)"')


def _esc(texto, atributo=False):
    texto = texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return texto.replace('"', "&quot;") if atributo else texto


def _unesc(texto):
    if "&" not in texto:
        return texto
    texto = _RE_ENT_NUM.sub(
        lambda m: chr(int(m.group(1)[1:], 16) if m.group(1)[0] in "xX" else int(m.group(1))), texto)
    return (texto.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
            .replace("&apos;", "'").replace("&amp;", "&"))


def _unesc_x(texto):
    """Excel guarda los caracteres de control como _x000D_."""
    return _RE_ESC_X.sub(lambda m: chr(int(m.group(1), 16)), texto) if "_x" in texto else texto


def _attrs(b):
    return {m.group(1): (m.group(2) if m.group(2) is not None else m.group(3)) for m in _RE_ATTR.finditer(b)}


_CACHE_IDX, _CACHE_LETRA = {}, {}


def _idx_col(letras):
    v = _CACHE_IDX.get(letras)
    if v is None:
        v = _CACHE_IDX[letras] = column_index_from_string(letras.decode("ascii"))
    return v


def _letra(col):
    v = _CACHE_LETRA.get(col)
    if v is None:
        v = _CACHE_LETRA[col] = get_column_letter(col).encode("ascii")
    return v


def _texto_t(b):
    """Texto de un fragmento con <t> (cadena compartida o inline), sin los fonéticos <rPh>."""
    if not b:
        return ""
    b = _RE_RPH.sub(b"", b)
    partes = [m.group(1) or b"" for m in _RE_T.finditer(b)]
    return _unesc_x(_unesc(b"".join(partes).decode("utf-8", "replace")))


def _leer_cadenas(datos):
    return [_texto_t(m.group(1)) for m in re.finditer(rb"<si\b[^>]*?(?:/>|>(.*?)</si>)", datos, re.S)]


def _texto_celda(cattrs, inner, cadenas):
    """Texto de una celda (None si no es texto o si es una fórmula)."""
    if not inner or b"<f" in inner:
        return None
    tipo = cattrs.get(b"t", b"n")
    if tipo == b"s":
        m = re.search(rb"<v>\s*(\d+)\s*</v>", inner)
        if m and int(m.group(1)) < len(cadenas):
            return cadenas[int(m.group(1))]
        return None
    if tipo == b"inlineStr":
        m = re.search(rb"<is\b[^>]*>(.*?)</is>", inner, re.S)
        return _texto_t(m.group(1)) if m else None
    if tipo == b"str":
        m = re.search(rb"<v>(.*?)</v>", inner, re.S)
        return _unesc_x(_unesc(m.group(1).decode("utf-8", "replace"))) if m else None
    return None


def _sub_attr(datos, nombre, fn, etiquetas=None):
    """Aplica fn(str) -> str al valor del atributo `nombre` (en cualquier etiqueta o solo en `etiquetas`)."""
    if etiquetas:
        patron = re.compile(rb"(<(?:" + b"|".join(etiquetas) + rb")\b[^>]*?\s" + nombre + rb'=")([^"]*)(")')
    else:
        patron = re.compile(rb"(\s" + nombre + rb'=")([^"]*)(")')

    def reemplazo(m):
        original = _unesc(m.group(2).decode("utf-8"))
        nuevo = fn(original)
        if nuevo == original:
            return m.group(0)
        return m.group(1) + _esc(nuevo, True).encode("utf-8") + m.group(3)

    return patron.sub(reemplazo, datos)


def _sub_texto(datos, patron, fn):
    """Patrón con tres grupos (apertura)(texto)(cierre): aplica fn(bytes) -> bytes al texto central."""
    return patron.sub(lambda m: m.group(1) + fn(m.group(2)) + m.group(3), datos)


# ---------------------------------------------------------------- plan de desplazamiento

class _Desp:
    """Qué se desplaza: n columnas a partir de la columna `ins` (base 1) de la hoja `titulo`."""

    def __init__(self, titulo, col_nombres, n=N_NUEVAS):
        self.titulo = titulo
        self.col = col_nombres
        self.ins = col_nombres + 1
        self.n = n

    def rangos(self, texto):
        return _desplazar_rangos(texto, self.titulo, self.ins, self.n)

    def formula(self, texto, propia=True):
        return _desplazar_formula("=" + texto, self.titulo, self.ins, self.n, propia)[1:]

    def formula_b(self, b, propia=True):
        """Formula en bytes (con escapes XML). Si no cambia, devuelve exactamente los mismos bytes."""
        original = _unesc(b.decode("utf-8"))
        nuevo = self.formula(original, propia)
        return b if nuevo == original else _esc(nuevo).encode("utf-8")


# ---------------------------------------------------------------- estructura del libro

def _resolver_ruta(base, destino):
    if destino.startswith("/"):
        return destino.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base), destino))


def _relaciones(zin, ruta_parte):
    """{Id: (Type, ruta_resuelta)} de las relaciones de una parte del paquete."""
    carpeta, nombre = posixpath.split(ruta_parte)
    ruta_rels = posixpath.join(carpeta, "_rels", nombre + ".rels")
    try:
        datos = zin.read(ruta_rels)
    except KeyError:
        return {}
    rels = {}
    for m in re.finditer(rb"<Relationship\b([^>]*?)/?>", datos):
        a = _attrs(m.group(1))
        if b"Id" in a and b"Target" in a and a.get(b"TargetMode") != b"External":
            rels[a[b"Id"].decode()] = (a.get(b"Type", b"").decode(), _resolver_ruta(ruta_parte, a[b"Target"].decode()))
    return rels


def _hojas(zin):
    """Hojas de cálculo del libro: lista de (nombre, ruta) y el índice de la hoja activa."""
    wb = zin.read("xl/workbook.xml")
    rels = _relaciones(zin, "xl/workbook.xml")
    hojas = []
    for m in re.finditer(rb"<sheet\b([^>]*?)/?>", wb):
        a = _attrs(m.group(1))
        rid = next((v for k, v in a.items() if k.endswith(b":id") or k == b"id"), None)
        if rid is None or rid.decode() not in rels:
            continue
        tipo, ruta = rels[rid.decode()]
        if tipo.endswith("/worksheet"):
            hojas.append((_unesc(a.get(b"name", b"").decode("utf-8")), ruta))
    if not hojas:
        raise ArchivoNoSoportado("El libro no contiene hojas de cálculo.")
    # Índice de la hoja activa entre TODAS las hojas (como openpyxl: activeTab)
    todas = [(_attrs(m.group(1))) for m in re.finditer(rb"<sheet\b([^>]*?)/?>", wb)]
    ma = re.search(rb"<workbookView\b[^>]*?\bactiveTab=\"(\d+)\"", wb)
    activa = int(ma.group(1)) if ma else 0
    ruta_activa = None
    if activa < len(todas):
        rid = next((v for k, v in todas[activa].items() if k.endswith(b":id") or k == b"id"), None)
        if rid is not None and rid.decode() in rels and rels[rid.decode()][0].endswith("/worksheet"):
            ruta_activa = rels[rid.decode()][1]
    if ruta_activa is None:
        ruta_activa = hojas[0][1]
    return hojas, ruta_activa


# ---------------------------------------------------------------- columnas (<cols>) y anchos

class _Anchos:
    """Ancho en píxeles de cada columna, para recalcular dibujos que cruzan el punto de inserción."""

    def __init__(self, definiciones, defecto_px=64):
        self.defs = definiciones          # [(min, max, ancho_en_caracteres | None, oculta)]
        self.defecto = defecto_px

    @staticmethod
    def _px(ancho):
        return int(((256 * ancho + int(128 / 7)) / 256) * 7)

    def px(self, col):
        for mn, mx, ancho, oculta in self.defs:
            if mn <= col <= mx:
                if oculta:
                    return 0
                return self._px(ancho) if ancho is not None else self.defecto
        return self.defecto


def _leer_cols(head):
    m = re.search(rb"<cols>(.*?)</cols>", head, re.S)
    cols = []
    if m:
        for c in re.finditer(rb"<col\b([^>]*?)/?>", m.group(1)):
            a = _attrs(c.group(1))
            cols.append(dict(a))
    return m, cols


def _construir_cols(cols_originales, col_nombres, n=N_NUEVAS):
    """Nuevas definiciones <col>: las de la derecha se desplazan y se agregan las columnas nuevas."""
    ins = col_nombres + 1
    resultado, estilo_nombres = [], None
    for a in cols_originales:
        mn, mx = int(a[b"min"]), int(a[b"max"])
        if mn <= col_nombres <= mx and b"style" in a:
            estilo_nombres = a[b"style"]
        if mx < ins:
            trozos = [(mn, mx)]
        elif mn >= ins:
            trozos = [(mn + n, mx + n)]
        else:                              # el rango atraviesa el punto de inserción: se parte en dos
            trozos = [(mn, ins - 1), (ins + n, mx + n)]
        for a1, b1 in trozos:
            if a1 > 16384:
                continue
            nuevo = dict(a)
            nuevo[b"min"], nuevo[b"max"] = str(a1).encode(), str(min(b1, 16384)).encode()
            resultado.append(nuevo)
    nuevas = {b"min": str(ins).encode(), b"max": str(ins + n - 1).encode(),
              b"width": str(ANCHO_NUEVAS).encode(), b"customWidth": b"1"}
    if estilo_nombres is not None:
        nuevas[b"style"] = estilo_nombres
    resultado.append(nuevas)
    resultado.sort(key=lambda a: int(a[b"min"]))
    return resultado


def _serializar_cols(cols):
    partes = []
    for a in cols:
        partes.append(b"<col " + b" ".join(k + b'="' + v + b'"' for k, v in a.items()) + b"/>")
    return b"<cols>" + b"".join(partes) + b"</cols>"


def _anchos_desde(cols, defecto_px):
    return _Anchos([(int(a[b"min"]), int(a[b"max"]),
                     float(a[b"width"]) if b"width" in a else None,
                     a.get(b"hidden") in (b"1", b"true")) for a in cols], defecto_px)


def _ancho_defecto_px(head):
    m = re.search(rb"<sheetFormatPr\b([^>]*?)/?>", head)
    if m:
        a = _attrs(m.group(1))
        if b"defaultColWidth" in a:
            return _Anchos._px(float(a[b"defaultColWidth"]))
    return 64


# ---------------------------------------------------------------- cabecera y cola de la hoja

def _ajustar_xml_hoja(b, p):
    """Referencias de todo lo que rodea a <sheetData>: selección, combinaciones, filtros, formatos
    condicionales, validaciones, hipervínculos, paneles..."""
    desp = p.rangos
    b = _sub_attr(b, b"sqref", desp)
    b = _sub_attr(b, b"activeCell", desp)

    def inicio_de_vista(texto):
        # Si el desplazamiento de la hoja empezaba justo donde se insertan las columnas, se deja ahí:
        # así las 4 columnas nuevas quedan a la vista al abrir el archivo (y no escondidas a la izquierda).
        m = _RE_CELDA.match(texto)
        if m and column_index_from_string(m.group(2).upper()) == p.ins:
            return texto
        return desp(texto)
    b = _sub_attr(b, b"topLeftCell", inicio_de_vista)
    b = _sub_attr(b, b"ref", desp, etiquetas=(b"mergeCell", b"sortState", b"sortCondition", b"hyperlink"))
    b = _sub_texto(b, re.compile(rb"(<formula[12]?>)(.*?)(</formula[12]?>)", re.S), p.formula_b)
    b = _sub_texto(b, re.compile(rb"(<xm:f>)(.*?)(</xm:f>)", re.S), p.formula_b)
    b = _sub_texto(b, re.compile(rb"(<xm:sqref>)(.*?)(</xm:sqref>)", re.S),
                   lambda t: p.rangos(_unesc(t.decode("utf-8"))).encode("utf-8"))

    # Paneles inmovilizados: si la inserción cae dentro de la zona fija, esta crece
    def panel(m):
        a = _attrs(m.group(1))
        if a.get(b"state", b"").startswith(b"frozen") and b"xSplit" in a:
            x = float(a[b"xSplit"])
            if x == int(x) and p.ins <= int(x):
                return m.group(0).replace(b'xSplit="' + a[b"xSplit"] + b'"', b'xSplit="%d"' % (int(x) + p.n))
        return m.group(0)
    b = re.sub(rb"<pane\b([^>]*?)/?>", panel, b)

    # Autofiltro de hoja (también se ensancha si terminaba justo en la columna de nombres)
    def autofiltro(m):
        a = _attrs(m.group(1))
        if b"ref" not in a:
            return m.group(0)
        c1, r1, c2, r2 = range_boundaries(a[b"ref"].decode())
        n1 = c1 + p.n if c1 >= p.ins else c1
        n2 = c2 + p.n if c2 >= p.col else c2
        nuevo_ref = f"{get_column_letter(n1)}{r1}:{get_column_letter(n2)}{r2}".encode()
        etiqueta = m.group(0)
        cuerpo = m.group(2)
        etiqueta = etiqueta.replace(b'ref="' + a[b"ref"] + b'"', b'ref="' + nuevo_ref + b'"', 1)
        if cuerpo and b"filterColumn" in cuerpo:
            def columna(mc):
                cid = int(mc.group(2))
                return mc.group(1) + (str(cid + p.n).encode() if c1 + cid >= p.ins else mc.group(2)) + mc.group(3)
            nuevo_cuerpo = re.sub(rb'(<filterColumn\b[^>]*?\bcolId=")(\d+)(")', columna, cuerpo)
            etiqueta = etiqueta.replace(cuerpo, nuevo_cuerpo, 1)
        return etiqueta
    b = re.sub(rb"<autoFilter\b([^>]*?)(?:/>|>(.*?)</autoFilter>)", autofiltro, b, flags=re.S)
    return b


def _ajustar_cabecera(head, p):
    """Cabecera de la hoja: dimensión, columnas (<cols>) y referencias."""
    head = _ajustar_xml_hoja(head, p)

    def dimension(m):
        ref = m.group(2).decode()
        try:
            c1, r1, c2, r2 = range_boundaries(ref)
        except Exception:
            return m.group(0)
        n1 = c1 + p.n if c1 >= p.ins else c1
        n2 = c2 + p.n if c2 >= p.ins else c2
        if c2 >= p.col:
            n2 = max(n2, p.ins + p.n - 1)
        return m.group(1) + f"{get_column_letter(n1)}{r1}:{get_column_letter(n2)}{r2}".encode() + m.group(3)
    head = re.sub(rb'(<dimension\b[^>]*?\sref=")([^"]*)(")', dimension, head)

    m, cols = _leer_cols(head)
    defecto = _ancho_defecto_px(head)
    anchos_viejos = _anchos_desde(cols, defecto)
    nuevas = _construir_cols(cols, p.col)
    anchos_nuevos = _anchos_desde(nuevas, defecto)
    bloque = _serializar_cols(nuevas)
    if m:
        head = head[:m.start()] + bloque + head[m.end():]
    else:
        head = head + bloque              # <cols> va justo antes de <sheetData>
    return head, anchos_viejos, anchos_nuevos


# ---------------------------------------------------------------- tablas de Excel

def _ajustar_tabla(xml, p):
    """Ajusta una tabla (xl/tables/tableN.xml). Devuelve (xml, info) con info = None si no cambia el
    encabezado, o (fila_encabezado, [4 nombres]) si la tabla contiene la columna de nombres."""
    mt = re.search(rb"<table\b([^>]*)>", xml)
    if not mt:
        return xml, None
    ta = _attrs(mt.group(1))
    if b"ref" not in ta:
        return xml, None
    c1, r1, c2, r2 = range_boundaries(ta[b"ref"].decode())
    if c2 < p.col:
        return xml, None                    # tabla totalmente a la izquierda: no cambia

    filas_enc = int(ta.get(b"headerRowCount", b"1"))
    filas_tot = int(ta.get(b"totalsRowCount", b"0"))
    contiene = c1 <= p.col
    n1, n2 = (c1, c2 + p.n) if contiene else (c1 + p.n, c2 + p.n)
    L = get_column_letter
    fila_datos_max = r2 - filas_tot
    ref_tabla = f"{L(n1)}{r1}:{L(n2)}{r2}"
    ref_filtro = f"{L(n1)}{r1}:{L(n2)}{fila_datos_max}"
    ref_orden = f"{L(n1)}{r1 + filas_enc}:{L(n2)}{fila_datos_max}"

    xml = xml.replace(b'ref="' + ta[b"ref"] + b'"', b'ref="' + ref_tabla.encode() + b'"', 1)

    def filtro(m):
        etiqueta, cuerpo = m.group(0), m.group(2)
        a = _attrs(m.group(1))
        if b"ref" in a:
            etiqueta = etiqueta.replace(b'ref="' + a[b"ref"] + b'"', b'ref="' + ref_filtro.encode() + b'"', 1)
        if contiene and cuerpo and b"filterColumn" in cuerpo:
            limite = p.col - c1 + 1

            def columna(mc):
                cid = int(mc.group(2))
                return mc.group(1) + (str(cid + p.n).encode() if cid >= limite else mc.group(2)) + mc.group(3)
            etiqueta = etiqueta.replace(
                cuerpo, re.sub(rb'(<filterColumn\b[^>]*?\bcolId=")(\d+)(")', columna, cuerpo), 1)
        return etiqueta
    xml = re.sub(rb"<autoFilter\b([^>]*?)(?:/>|>(.*?)</autoFilter>)", filtro, xml, count=1, flags=re.S)

    def orden(m):
        a = _attrs(m.group(1))
        etiqueta = m.group(0)
        if b"ref" in a:
            etiqueta = etiqueta.replace(b'ref="' + a[b"ref"] + b'"', b'ref="' + ref_orden.encode() + b'"', 1)
        return etiqueta
    xml = re.sub(rb"<sortState\b([^>]*?)(?:/>|>.*?</sortState>)",
                 lambda m: _sub_attr(orden(m), b"ref", p.rangos, etiquetas=(b"sortCondition",)), xml, flags=re.S)

    nuevo_info = None
    mc = re.search(rb"<tableColumns\b([^>]*)>(.*?)</tableColumns>", xml, re.S)
    if mc:
        columnas = [m.group(0) for m in re.finditer(rb"<tableColumn\b[^>]*?(?:/>|>.*?</tableColumn>)", mc.group(2), re.S)]
        columnas = [_sub_texto(c, re.compile(rb"(<calculatedColumnFormula\b[^>]*>)(.*?)(</calculatedColumnFormula>)", re.S),
                               p.formula_b) for c in columnas]
        if contiene:
            usados = {_unesc(_attrs(c)[b"name"].decode("utf-8")).lower() for c in columnas if b"name" in _attrs(c)}
            ids = [int(_attrs(c).get(b"id", b"0")) for c in columnas]
            siguiente = max(ids + [0]) + 1
            nombres = []
            for h in HEADERS_NUEVOS:
                nombre, k = h, 1
                while nombre.lower() in usados:      # Excel exige nombres únicos, sin distinguir mayúsculas
                    nombre = f"{h}_{k}"
                    k += 1
                usados.add(nombre.lower())
                nombres.append(nombre)
            nuevas = [b'<tableColumn id="%d" name="%b"/>' % (siguiente + i, _esc(nm, True).encode("utf-8"))
                      for i, nm in enumerate(nombres)]
            pos = p.col - c1 + 1
            columnas = columnas[:pos] + nuevas + columnas[pos:]
            if filas_enc > 0:
                nuevo_info = (r1, nombres)
        abre = re.sub(rb'\bcount="\d+"', b'count="%d"' % len(columnas), b"<tableColumns" + mc.group(1) + b">")
        xml = xml[:mc.start()] + abre + b"".join(columnas) + b"</tableColumns>" + xml[mc.end():]
    return xml, nuevo_info


# ---------------------------------------------------------------- dibujos, comentarios, gráficos

def _ajustar_dibujo(xml, p, anchos_viejos, anchos_nuevos):
    """Mueve los objetos (imágenes, logos, formas, gráficos) que quedan a la derecha del punto de
    inserción. Los que lo cruzan se estiran o conservan su tamaño según su propiedad editAs."""
    i0 = p.ins - 1                                   # índice base 0 de la primera columna desplazada
    re_col = re.compile(rb"(<(?:\w+:)?col>)(\d+)(</(?:\w+:)?col>)")
    re_off = re.compile(rb"(<(?:\w+:)?colOff>)(-?\d+)(</(?:\w+:)?colOff>)")
    re_punto = re.compile(rb"(<(?:\w+:)?(from|to)>)(.*?)(</(?:\w+:)?\2>)", re.S)

    def leer(cuerpo):
        c = int(re_col.search(cuerpo).group(2))
        mo = re_off.search(cuerpo)
        return c, int(mo.group(2)) if mo else 0

    def escribir(cuerpo, col, off):
        cuerpo = re_col.sub(lambda m: m.group(1) + str(col).encode() + m.group(3), cuerpo, count=1)
        return re_off.sub(lambda m: m.group(1) + str(off).encode() + m.group(3), cuerpo, count=1)

    def ancla(m):
        prefijo, tipo, attrs, cuerpo = m.group(1), m.group(2), m.group(3), m.group(4)
        puntos = {mm.group(2): mm for mm in re_punto.finditer(cuerpo)}
        if b"from" not in puntos:
            return m.group(0)
        f_col, f_off = leer(puntos[b"from"].group(3))
        cambios = {}                                   # b"from" / b"to" -> nuevo contenido
        if tipo == b"oneCellAnchor" or b"to" not in puntos:
            if f_col >= i0:
                cambios[b"from"] = escribir(puntos[b"from"].group(3), f_col + p.n, f_off)
        else:
            t_col, t_off = leer(puntos[b"to"].group(3))
            edita = _attrs(attrs).get(b"editAs", b"twoCell")
            if f_col >= i0:                           # entero a la derecha: se mueve
                cambios[b"from"] = escribir(puntos[b"from"].group(3), f_col + p.n, f_off)
                cambios[b"to"] = escribir(puntos[b"to"].group(3), t_col + p.n, t_off)
            elif t_col >= i0:                         # cruza el punto de inserción
                if edita == b"twoCell":               # mover y cambiar de tamaño: se estira
                    cambios[b"to"] = escribir(puntos[b"to"].group(3), t_col + p.n, t_off)
                else:                                 # conserva su tamaño en píxeles
                    emu = sum(anchos_viejos.px(c + 1) * 9525 for c in range(f_col, t_col)) - f_off + t_off
                    col, pos, resto = f_col, f_off, emu
                    while True:
                        ancho = anchos_nuevos.px(col + 1) * 9525
                        if ancho <= 0 or pos + resto <= ancho or col > 16383:
                            break
                        resto -= ancho - pos
                        pos, col = 0, col + 1
                    cambios[b"to"] = escribir(puntos[b"to"].group(3), col, pos + resto)
        if not cambios:
            return m.group(0)
        nuevo = cuerpo
        for clave in sorted(cambios, key=lambda k: -puntos[k].start(3)):      # de atrás hacia adelante
            mm = puntos[clave]
            nuevo = nuevo[:mm.start(3)] + cambios[clave] + nuevo[mm.end(3):]
        return b"<" + prefijo + tipo + attrs + b">" + nuevo + b"</" + prefijo + tipo + b">"

    return re.sub(rb"<((?:\w+:)?)(twoCellAnchor|oneCellAnchor)\b([^>]*)>(.*?)</\1\2>", ancla, xml, flags=re.S)


def _ajustar_vml(xml, p):
    i0 = p.ins - 1

    def columna(m):
        v = int(m.group(2))
        return m.group(1) + str(v + p.n if v >= i0 else v).encode() + m.group(3)

    def ancla(m):
        partes = [x.strip() for x in m.group(2).split(b",")]
        for i in (0, 4):
            if len(partes) > i and partes[i].isdigit() and int(partes[i]) >= i0:
                partes[i] = str(int(partes[i]) + p.n).encode()
        return m.group(1) + b", ".join(partes) + m.group(3)

    xml = re.sub(rb"(<x:Column>)(\d+)(</x:Column>)", columna, xml)
    return re.sub(rb"(<x:Anchor>)(.*?)(</x:Anchor>)", ancla, xml, flags=re.S)


# ---------------------------------------------------------------- celdas nuevas y reescritura de filas

def _celda_nueva(col, fila, texto, estilo):
    ref = _letra(col) + str(fila).encode()
    s = b' s="' + estilo + b'"' if estilo and estilo != b"0" else b""
    if not texto:
        return b'<c r="' + ref + b'"' + s + b"/>"
    conservar = b' xml:space="preserve"' if (texto != texto.strip() or "  " in texto) else b""
    return (b'<c r="' + ref + b'"' + s + b' t="inlineStr"><is><t' + conservar + b">"
            + _esc(texto).encode("utf-8") + b"</t></is></c>")


class _Plan(_Desp):
    """Todo lo necesario para reescribir las filas de la hoja."""

    def __init__(self, titulo, col_nombres, fila_header):
        super().__init__(titulo, col_nombres)
        self.fila_header = fila_header
        self.cabeceras = {fila_header: list(HEADERS_NUEVOS)}   # fila -> 4 títulos
        self.resultados = {}                                   # fila -> Resultado
        self.expandir = set()                                  # grupos de fórmulas compartidas a expandir
        self.maestras = {}                                     # si -> (coordenada, fórmula)


def _formula_celda(inner, fila, col, p):
    """Reescribe el <f> de una celda: desplaza sus referencias y expande las fórmulas compartidas que
    quedarían partidas por la inserción (las de varias columnas)."""
    def f(m):
        a = _attrs(m.group(1))
        texto = m.group(2)
        if a.get(b"t") == b"shared" and a.get(b"si") in p.expandir:
            if texto is not None:                                   # la maestra pasa a fórmula normal
                return b"<f>" + p.formula_b(texto) + b"</f>"
            coord, base = p.maestras[a[b"si"]]                      # un hijo: se traduce desde la maestra
            nueva = Translator("=" + base, origin=coord).translate_formula(f"{get_column_letter(col)}{fila}")[1:]
            return b"<f>" + _esc(p.formula(nueva)).encode("utf-8") + b"</f>"
        if texto is None:
            return m.group(0)                                       # hijo de un grupo compartido: no cambia
        atributos = m.group(1)
        if b"ref" in a:                                             # maestra compartida o fórmula matricial
            nuevo_ref = p.rangos(a[b"ref"].decode()).encode()
            atributos = atributos.replace(b'ref="' + a[b"ref"] + b'"', b'ref="' + nuevo_ref + b'"', 1)
        return b"<f" + atributos + b">" + p.formula_b(texto) + b"</f>"
    return _RE_F.sub(f, inner)


def _reescribir_fila(fila_attrs, inner, fila, p):
    """Desplaza las celdas de la derecha, inserta las 4 nuevas junto a la de nombres y devuelve la fila."""
    celdas = list(_RE_CEL.finditer(inner))
    resto = _RE_CEL.sub(b"", inner).strip()           # contenido que no son celdas (extLst...), se conserva
    partes, insertadas, estilo_nombre, hay_nombre = [], False, None, False
    siguiente = 1
    ins, n = p.ins, p.n
    for cm in celdas:
        cattrs, cinner = cm.group(1), cm.group(2)
        mr = _RE_R_CELDA.search(cattrs)
        col = _idx_col(mr.group(1)) if mr else siguiente
        siguiente = col + 1
        if col == p.col:
            hay_nombre = True
            ms = _RE_S_ATTR.search(cattrs)
            estilo_nombre = ms.group(1) if ms else None
        if col >= ins and not insertadas:
            if hay_nombre:
                partes.append(_celdas_para(fila, estilo_nombre, p))
            insertadas = True
        if col >= ins:
            cattrs = _RE_R_ATTR.sub(lambda m: m.group(1) + b'r="' + _letra(col + n) + str(fila).encode() + b'"', cattrs, count=1)
        if cinner is not None and b"<f" in cinner:
            cinner = _formula_celda(cinner, fila, col, p)
        if cinner is None:
            partes.append(b"<c" + cattrs + b"/>")
        else:
            partes.append(b"<c" + cattrs + b">" + cinner + b"</c>")
    if hay_nombre and not insertadas:
        partes.append(_celdas_para(fila, estilo_nombre, p))
    fila_attrs = _RE_SPANS.sub(b"", fila_attrs)         # 'spans' es solo una pista de optimización
    return b"<row" + fila_attrs + b">" + b"".join(partes) + resto + b"</row>"


def _celdas_para(fila, estilo, p):
    if fila in p.cabeceras:
        textos = p.cabeceras[fila]
    elif fila in p.resultados:
        r = p.resultados[fila]
        textos = [r.p_ape, r.s_ape, r.p_nom, r.s_nom]
    else:
        textos = None
    if textos is None and (not estilo or estilo == b"0"):
        return b""                                       # nada que escribir ni formato que continuar
    return b"".join(_celda_nueva(p.ins + i, fila, textos[i] if textos else "", estilo) for i in range(p.n))


# ---------------------------------------------------------------- recorrido de la hoja

def _numero_fila(fattrs, anterior):
    r = fattrs.get(b"r")
    return int(r) if r else anterior + 1


def _buscar_encabezado(hoja, ini, fin, cadenas):
    """Primera celda con 'NOMBRE' o 'APELLIDO' por encima de FILA_MAX_ENCABEZADO (lectura por filas,
    de izquierda a derecha, igual que la versión anterior). Devuelve (fila, columna, textos_de_la_fila)."""
    anterior = 0
    for fm in _RE_FILA.finditer(hoja, ini, fin):
        fila = _numero_fila(_attrs(fm.group(1)), anterior)
        anterior = fila
        if fila >= FILA_MAX_ENCABEZADO:
            break
        if fm.group(2) is None:
            continue
        textos, hallada, siguiente = [], None, 1
        for cm in _RE_CEL.finditer(fm.group(2)):
            ca = _attrs(cm.group(1))
            mr = _RE_REF_B.match(ca.get(b"r", b""))
            col = _idx_col(mr.group(1)) if mr else siguiente
            siguiente = col + 1
            val = (_texto_celda(ca, cm.group(2), cadenas) or "").upper()
            textos.append(val)
            if hallada is None and ("NOMBRE" in val or "APELLIDO" in val) and val not in HEADERS_NUEVOS:
                hallada = col
        if hallada is not None:
            return fila, hallada, textos
    return None


def _recoger_nombres(hoja, ini, fin, cadenas, p, filas_omitir):
    """Textos de la columna de nombres debajo del encabezado y grupos de fórmulas compartidas."""
    nombres, anterior = [], 0
    hay_compartidas = hoja.find(b't="shared"', ini, fin) != -1
    multicolumna = set()
    for fm in _RE_FILA.finditer(hoja, ini, fin):
        fila = _numero_fila(_attrs(fm.group(1)), anterior)
        anterior = fila
        if fm.group(2) is None:
            continue
        inner = fm.group(2)
        buscar_nombre = fila > p.fila_header and fila not in filas_omitir
        if not buscar_nombre and not (hay_compartidas and b't="shared"' in inner):
            continue
        siguiente = 1
        for cm in _RE_CEL.finditer(inner):
            ca = _attrs(cm.group(1))
            mr = _RE_REF_B.match(ca.get(b"r", b""))
            col = _idx_col(mr.group(1)) if mr else siguiente
            siguiente = col + 1
            cinner = cm.group(2)
            if buscar_nombre and col == p.col:
                texto = _texto_celda(ca, cinner, cadenas)
                if texto is not None and texto.strip():
                    nombres.append((fila, texto))
            if hay_compartidas and cinner and b't="shared"' in cinner:
                for fmm in _RE_F.finditer(cinner):
                    fa = _attrs(fmm.group(1))
                    if fa.get(b"t") == b"shared" and b"ref" in fa and fmm.group(2) is not None:
                        c1, _, c2, _ = range_boundaries(fa[b"ref"].decode())
                        si = fa[b"si"]
                        p.maestras[si] = (f"{get_column_letter(col)}{fila}", _unesc(fmm.group(2).decode("utf-8")))
                        if c1 != c2:
                            multicolumna.add(si)
            if col > p.col and not hay_compartidas:
                break
    p.expandir = multicolumna
    return nombres


def _filas_de_totales(tablas_xml):
    omitir = set()
    for xml in tablas_xml:
        mt = re.search(rb"<table\b([^>]*)>", xml)
        if not mt:
            continue
        ta = _attrs(mt.group(1))
        tot = int(ta.get(b"totalsRowCount", b"0"))
        if tot and b"ref" in ta:
            _, _, _, r2 = range_boundaries(ta[b"ref"].decode())
            omitir.update(range(r2 - tot + 1, r2 + 1))
    return omitir


def _escribir_filas(hoja, ini, fin, p, destino):
    """Reescribe las filas existentes y las escribe por tandas (memoria acotada)."""
    lote, tam, anterior = [], 0, 0
    for fm in _RE_FILA.finditer(hoja, ini, fin):
        fila = _numero_fila(_attrs(fm.group(1)), anterior)
        anterior = fila
        if fm.group(2) is None:
            fragmento = b"<row" + _RE_SPANS.sub(b"", fm.group(1)) + b"/>"
        else:
            fragmento = _reescribir_fila(fm.group(1), fm.group(2), fila, p)
        lote.append(fragmento)
        tam += len(fragmento)
        if tam > (1 << 20):
            destino.write(b"".join(lote))
            lote, tam = [], 0
    if lote:
        destino.write(b"".join(lote))


# ---------------------------------------------------------------- procesamiento del paquete

def _copiar_info(info):
    nueva = zipfile.ZipInfo(info.filename, date_time=info.date_time)
    nueva.compress_type = info.compress_type
    nueva.external_attr = info.external_attr
    nueva.create_system = info.create_system
    return nueva


def procesar_xlsx(datos):
    """Separa los nombres de un .xlsx editando su XML. Devuelve (bytes_del_nuevo_xlsx, estadísticas).

    No carga el libro en memoria: solo recorre las filas que existen y copia byte a byte todo lo que
    no cambia (imágenes, logos, tema, estilos...)."""
    try:
        zin = zipfile.ZipFile(io.BytesIO(datos))
    except zipfile.BadZipFile:
        raise ArchivoNoSoportado("El archivo no es un .xlsx válido (o está dañado).")
    infos = zin.infolist()
    if sum(i.file_size for i in infos) > LIMITE_DESCOMPRIMIDO_MB * 1024 * 1024:
        raise ArchivoNoSoportado(
            f"El contenido del archivo supera {LIMITE_DESCOMPRIMIDO_MB} MB descomprimido; es demasiado pesado para procesarlo aquí.")
    existentes = {i.filename for i in infos}
    if "xl/workbook.xml" not in existentes:
        raise ArchivoNoSoportado("El archivo no parece un libro de Excel (.xlsx).")

    hojas, ruta_hoja = _hojas(zin)
    titulo = next(nombre for nombre, ruta in hojas if ruta == ruta_hoja)
    hoja = zin.read(ruta_hoja)
    cadenas = _leer_cadenas(zin.read("xl/sharedStrings.xml")) if "xl/sharedStrings.xml" in existentes else []

    ini_tag = hoja.find(b"<sheetData")
    if ini_tag == -1:
        raise ArchivoNoSoportado("La hoja tiene un formato interno que no se puede procesar. Guárdala de nuevo desde Excel.")
    fin_tag = hoja.find(b">", ini_tag)
    fin_datos = hoja.rfind(b"</sheetData>")
    if hoja[fin_tag - 1:fin_tag] == b"/" or fin_datos == -1:
        raise ColumnaNoEncontrada("La hoja no tiene datos.")
    ini_datos = fin_tag + 1

    encabezado = _buscar_encabezado(hoja, ini_datos, fin_datos, cadenas)
    if encabezado is None:
        raise ColumnaNoEncontrada("No se encontró una columna válida de Nombres.")
    fila_header, col_nombres, textos_enc = encabezado
    p = _Plan(titulo, col_nombres, fila_header)

    # Partes relacionadas con la hoja
    rels = _relaciones(zin, ruta_hoja)
    ruta_tablas = [r for t, r in rels.values() if t.endswith("/table") and r in existentes]
    tablas = {r: zin.read(r) for r in ruta_tablas}
    filas_omitir = _filas_de_totales(tablas.values())

    nombres = _recoger_nombres(hoja, ini_datos, fin_datos, cadenas, p, filas_omitir)
    resultados = resolver_lista([t for _, t in nombres])
    p.resultados = {fila: res for (fila, _), res in zip(nombres, resultados)}

    nuevas_partes = {}
    for ruta, xml in tablas.items():
        nuevo, info = _ajustar_tabla(xml, p)
        nuevas_partes[ruta] = nuevo
        if info:
            p.cabeceras[info[0]] = info[1]

    # Cabecera y cola de la hoja
    cabecera, anchos_viejos, anchos_nuevos = _ajustar_cabecera(hoja[:ini_tag], p)
    cola = _ajustar_xml_hoja(hoja[fin_datos + len(b"</sheetData>"):], p)

    # Dibujos, comentarios y gráficos
    for tipo, ruta in rels.values():
        if ruta not in existentes:
            continue
        if tipo.endswith("/drawing"):
            nuevas_partes[ruta] = _ajustar_dibujo(zin.read(ruta), p, anchos_viejos, anchos_nuevos)
        elif tipo.endswith("/comments") or tipo.endswith("/threadedComment"):
            nuevas_partes[ruta] = _sub_attr(zin.read(ruta), b"ref", p.rangos, etiquetas=(b"comment", b"threadedComment"))
        elif tipo.endswith("/vmlDrawing"):
            nuevas_partes[ruta] = _ajustar_vml(zin.read(ruta), p)

    # Nombres definidos y otras hojas que apuntan a esta
    wb = zin.read("xl/workbook.xml")
    nuevo_wb = _sub_texto(wb, re.compile(rb"(<definedName\b[^>]*>)(.*?)(</definedName>)", re.S),
                          lambda t: p.formula_b(t, propia=False))
    if nuevo_wb != wb:
        nuevas_partes["xl/workbook.xml"] = nuevo_wb
    marca = titulo.encode("utf-8")
    for nombre, ruta in hojas:
        if ruta == ruta_hoja or ruta not in existentes:
            continue
        otra = zin.read(ruta)
        if marca in otra and b"<f" in otra:
            nueva = _sub_texto(otra, re.compile(rb"(<f\b[^>]*>)(.*?)(</f>)", re.S), lambda t: p.formula_b(t, propia=False))
            if nueva != otra:
                nuevas_partes[ruta] = nueva
    for nombre in existentes:
        if re.match(r"xl/charts/chart\d+\.xml$", nombre):
            g = zin.read(nombre)
            if marca in g:
                nuevo = _sub_texto(g, re.compile(rb"(<c:f>)(.*?)(</c:f>)", re.S), lambda t: p.formula_b(t, propia=False))
                if nuevo != g:
                    nuevas_partes[nombre] = nuevo
        elif re.match(r"xl/pivotCache/pivotCacheDefinition\d+\.xml$", nombre):
            g = zin.read(nombre)
            nuevo = re.sub(rb"<worksheetSource\b[^>]*?/?>",
                           lambda m: (_sub_attr(m.group(0), b"ref", p.rangos)
                                      if _unesc(_attrs(m.group(0)).get(b"sheet", b"").decode("utf-8")) == titulo else m.group(0)), g)
            if nuevo != g:
                nuevas_partes[nombre] = nuevo

    # La cadena de cálculo quedaría desactualizada: se elimina y Excel la reconstruye al abrir
    quitar = set()
    if "xl/calcChain.xml" in existentes:
        quitar.add("xl/calcChain.xml")
        ct = zin.read("[Content_Types].xml")
        nuevas_partes["[Content_Types].xml"] = re.sub(rb'<Override\b[^>]*calcChain[^>]*?/>', b"", ct)
        rw = zin.read("xl/_rels/workbook.xml.rels")
        nuevas_partes["xl/_rels/workbook.xml.rels"] = re.sub(rb'<Relationship\b[^>]*calcChain[^>]*?/>', b"", rw)

    # Escritura del nuevo paquete: lo que no cambia se copia sin tocar
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in infos:
            if info.filename in quitar:
                continue
            nueva_info = _copiar_info(info)
            if info.filename == ruta_hoja:
                with zout.open(nueva_info, "w") as destino:
                    destino.write(cabecera + hoja[ini_tag:ini_datos])
                    _escribir_filas(hoja, ini_datos, fin_datos, p, destino)
                    destino.write(b"</sheetData>" + cola)
            elif info.filename in nuevas_partes:
                zout.writestr(nueva_info, nuevas_partes[info.filename])
            elif info.is_dir():
                zout.writestr(nueva_info, b"")
            else:
                with zin.open(info) as origen, zout.open(nueva_info, "w") as destino:
                    shutil.copyfileobj(origen, destino, 1 << 20)

    # Estadísticas y avisos
    conteo = {"A": 0, "B": 0, "AUTO": 0}
    vista_previa, por_revisar = [], []
    for (fila, texto), res in zip(nombres, resultados):
        conteo[res.origen] += 1
        resumen = {"Fila": fila, "Original": texto,
                   "PRIMER APELLIDO": res.p_ape, "SEGUNDO APELLIDO": res.s_ape,
                   "PRIMER NOMBRE": res.p_nom, "SEGUNDO NOMBRE": res.s_nom}
        if len(vista_previa) < 10:
            vista_previa.append(resumen)
        if res.revisar:
            por_revisar.append({**resumen, "Criterio": res.motivo})

    advertencias = []
    if any(h in textos_enc for h in HEADERS_NUEVOS):
        advertencias.append(
            "La hoja ya tenía columnas con los títulos PRIMER APELLIDO / SEGUNDO APELLIDO / PRIMER NOMBRE / "
            "SEGUNDO NOMBRE. No se tocaron: las nuevas se agregaron junto a la columna de nombres, así que "
            "esos títulos aparecen dos veces.")
    sin_ajustar = sorted({etiqueta for prefijo, etiqueta in (
        ("xl/pivotTables/", "tablas dinámicas"), ("xl/slicers/", "segmentaciones"),
        ("xl/timelines/", "escalas de tiempo"), ("xl/queryTables/", "consultas de datos externos"),
        ("xl/ctrlProps/", "controles de formulario")) if any(n.startswith(prefijo) for n in existentes)})
    if sin_ajustar:
        advertencias.append(
            "El archivo contiene " + ", ".join(sin_ajustar) + ". Se conservaron, pero sus referencias a celdas "
            "no se ajustaron; si alguna queda a la derecha de la columna de nombres, revísala.")

    return salida.getvalue(), {"registros": len(nombres), "conteo": conteo, "vista_previa": vista_previa,
                               "revisar": por_revisar, "advertencias": advertencias}


# =============================================================================
# 5. ESTILOS (solo estética)
# =============================================================================

NOMBRE_APP = "Desglosa"
LEMA_APP = "Separador inteligente de nombres"
_ICONO_ARCHIVO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "favicon.png")
ICONO_APP = _ICONO_ARCHIVO if os.path.exists(_ICONO_ARCHIVO) else "📊"

LOGO_SVG = """
<svg class="logo" viewBox="0 0 64 64" xmlns="http://www.w3.org/2000/svg" aria-hidden="true">
  <defs>
    <linearGradient id="g" x1="0" y1="0" x2="1" y2="1">
      <stop offset="0" stop-color="#1d4ed8"/><stop offset="1" stop-color="#10b981"/>
    </linearGradient>
  </defs>
  <rect width="64" height="64" rx="14" fill="url(#g)"/>
  <rect x="13" y="15" width="38" height="9" rx="4.5" fill="#fff"/>
  <rect x="30.5" y="27" width="3" height="4.5" fill="#fff" fill-opacity=".75"/>
  <path d="M27 31h10l-5 5.5z" fill="#fff" fill-opacity=".75"/>
  <rect x="13" y="39" width="7.4" height="10" rx="2.2" fill="#fff"/>
  <rect x="23" y="39" width="7.4" height="10" rx="2.2" fill="#fff"/>
  <rect x="33" y="39" width="7.4" height="10" rx="2.2" fill="#fff"/>
  <rect x="43" y="39" width="7.4" height="10" rx="2.2" fill="#fff"/>
</svg>
"""

# Paletas de color: el resto del CSS solo usa estas variables.
# Oscuro es el tema por defecto; Claro se activa solo con el botón.
TEMA_OSCURO = """
    color-scheme: dark;
    --tinta: #e2e8f0;
    --gris: #94a3b8;
    --borde: #263449;
    --fondo-app: #0b1220;
    --fondo-card: #111a2e;
    --fondo-suave: #0e1627;
    --acento-texto: #7db1ff;
    --punteado: #475569;
    --glow-azul: rgba(59,130,246,.16);
    --glow-verde: rgba(16,185,129,.12);
    --sombra: rgba(0,0,0,.65);
"""

TEMA_CLARO = """
    color-scheme: light;
    --tinta: #0f172a;
    --gris: #64748b;
    --borde: #e2e8f0;
    --fondo-app: #f8fafc;
    --fondo-card: #ffffff;
    --fondo-suave: #f8fafc;
    --acento-texto: #1d4ed8;
    --punteado: #94a3b8;
    --glow-azul: rgba(59,130,246,.10);
    --glow-verde: rgba(16,185,129,.10);
    --sombra: rgba(15,23,42,.25);
"""

CSS_TEMA_OSCURO = f"<style>:root {{ {TEMA_OSCURO} }}</style>"
CSS_TEMA_CLARO = f"<style>:root {{ {TEMA_CLARO} }}</style>"

ESTILOS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

:root {
    --verde: #10b981;
    --verde-osc: #059669;
    --azul: #3b82f6;
    --azul-osc: #1d4ed8;
}

html, body, [class*="css"], .stApp {
    font-family: 'Inter', sans-serif;
}

/* Fondo general */
.stApp {
    background:
        radial-gradient(1000px 500px at 10% -10%, var(--glow-azul), transparent 60%),
        radial-gradient(900px 500px at 100% 0%, var(--glow-verde), transparent 60%),
        var(--fondo-app) !important;
    transition: background-color .25s ease;
}
[data-testid="stAppViewContainer"], [data-testid="stMain"] { background: transparent !important; }

/* Ocultar elementos por defecto de Streamlit */
#MainMenu, footer, header[data-testid="stHeader"] { visibility: hidden; height: 0; }
.block-container { max-width: 780px; padding-top: 1.4rem; padding-bottom: 3rem; }

/* Botón discreto de tema (esquina inferior izquierda) */
.st-key-selector_tema {
    position: fixed; left: 16px; bottom: 16px; z-index: 1000;
    width: auto !important;
}
.st-key-selector_tema button {
    width: 38px; height: 38px; min-height: 0; padding: 0;
    border-radius: 50%;
    background: var(--fondo-card);
    border: 1px solid var(--borde);
    color: var(--gris);
    opacity: .55;
    box-shadow: 0 6px 16px -8px var(--sombra);
    transition: opacity .2s ease, border-color .2s ease, transform .2s ease;
}
.st-key-selector_tema button p { color: inherit; font-size: 1.1rem; line-height: 1; }
.st-key-selector_tema button:hover {
    opacity: 1; color: var(--tinta); border-color: var(--azul); transform: scale(1.06);
}

/* HERO */
.hero {
    background: linear-gradient(135deg, #0f172a 0%, #1e3a8a 55%, #059669 120%);
    border: 1px solid rgba(255,255,255,.06);
    border-radius: 24px;
    padding: 2.3rem 2rem 2.2rem 2rem;
    color: #fff;
    position: relative;
    overflow: hidden;
    box-shadow: 0 20px 40px -18px var(--sombra);
    margin-bottom: 1.4rem;
}
.hero::after {
    content: "";
    position: absolute; right: -60px; top: -60px;
    width: 220px; height: 220px; border-radius: 50%;
    background: rgba(255,255,255,.08);
}
.hero::before {
    content: "";
    position: absolute; right: 60px; bottom: -90px;
    width: 180px; height: 180px; border-radius: 50%;
    background: rgba(16,185,129,.25);
}
.hero .badge {
    display: inline-block;
    font-size: .72rem; font-weight: 600; letter-spacing: .08em;
    text-transform: uppercase;
    background: rgba(255,255,255,.14);
    border: 1px solid rgba(255,255,255,.25);
    padding: .3rem .7rem; border-radius: 999px; margin-bottom: 1rem;
    position: relative; z-index: 1;
}
.hero .marca { display: flex; align-items: center; gap: .85rem; position: relative; z-index: 1; }
.hero .logo { width: 52px; height: 52px; flex: none; filter: drop-shadow(0 8px 14px rgba(0,0,0,.35)); }
.hero h1 {
    font-size: 2.5rem; font-weight: 800; line-height: 1.1; letter-spacing: -.02em;
    margin: 0; padding: 0; color: #fff;
}
.hero .lema {
    font-size: 1.02rem; font-weight: 600; color: rgba(255,255,255,.92);
    margin: .55rem 0 .6rem 0; position: relative; z-index: 1;
}
.hero p {
    font-size: .98rem; color: rgba(255,255,255,.78);
    margin: 0; max-width: 560px; line-height: 1.55; position: relative; z-index: 1;
}

/* Chips de características */
.chips { display: flex; flex-wrap: wrap; gap: .5rem; margin-top: 1.1rem; position: relative; z-index: 1; }
.chip {
    font-size: .78rem; font-weight: 500; color: #fff;
    background: rgba(255,255,255,.12);
    border: 1px solid rgba(255,255,255,.2);
    padding: .3rem .75rem; border-radius: 999px;
}

/* Encabezados de paso */
.paso {
    display: flex; align-items: center; gap: .7rem;
    margin: 1.8rem 0 .7rem 0;
}
.paso .num {
    width: 30px; height: 30px; border-radius: 50%;
    background: linear-gradient(135deg, var(--azul), var(--verde));
    color: #fff; font-weight: 700; font-size: .9rem;
    display: flex; align-items: center; justify-content: center;
    box-shadow: 0 6px 14px -4px rgba(59,130,246,.55);
}
.paso .titulo { font-size: 1.05rem; font-weight: 700; color: var(--tinta); }
.paso .sub { font-size: .85rem; color: var(--gris); font-weight: 400; }
.pregunta { margin: 1.3rem 0 .55rem 0; font-size: .9rem; font-weight: 600; color: var(--gris); }

/* Zona de carga de archivo */
[data-testid="stFileUploader"] section {
    display: flex; flex-direction: column; align-items: center; justify-content: center;
    gap: 1.1rem;
    text-align: center;
    padding: 2.2rem 1.2rem 2rem 1.2rem;
    background:
        radial-gradient(420px 160px at 50% 0%, rgba(59,130,246,.10), transparent 70%),
        var(--fondo-card);
    border: 2px dashed var(--punteado);
    border-radius: 22px;
    box-shadow: 0 14px 30px -22px var(--sombra);
    transition: all .2s ease;
}
[data-testid="stFileUploader"] section:hover {
    border-color: var(--verde);
    background:
        radial-gradient(420px 160px at 50% 0%, rgba(16,185,129,.14), transparent 70%),
        var(--fondo-card);
    transform: translateY(-2px);
    box-shadow: 0 20px 36px -22px rgba(16,185,129,.45);
}
/* icono */
[data-testid="stFileUploaderDropzoneInstructions"] {
    display: flex; flex-direction: column; align-items: center; gap: .8rem;
}
[data-testid="stFileUploaderDropzoneInstructions"] svg,
[data-testid="stFileUploaderDropzoneInstructions"] [data-testid="stIconMaterial"] {
    box-sizing: content-box;
    width: 1.9rem; height: 1.9rem; font-size: 1.9rem; line-height: 1.9rem;
    padding: .85rem;
    color: #fff; fill: #fff;
    overflow: hidden; white-space: nowrap;
    border-radius: 18px;
    background: linear-gradient(135deg, var(--azul), var(--verde));
    box-shadow: 0 12px 22px -10px rgba(59,130,246,.65);
}
/* textos en español (reemplazan los de Streamlit) */
[data-testid="stFileUploaderDropzoneInstructions"] > div > * { display: none; }
[data-testid="stFileUploaderDropzoneInstructions"] > div::before {
    content: "Arrastra tu archivo Excel aquí";
    display: block;
    font-size: 1.02rem; font-weight: 700; color: var(--tinta);
}
[data-testid="stFileUploaderDropzoneInstructions"] > div::after {
    content: "o selecciónalo desde tu equipo · solo .xlsx";
    display: block; margin-top: .3rem;
    font-size: .82rem; font-weight: 400; color: var(--gris);
}
/* botón */
[data-testid="stFileUploader"] section button {
    font-size: 0 !important;
    padding: .72rem 1.5rem;
    border: none;
    border-radius: 999px;
    background: linear-gradient(135deg, var(--azul) 0%, var(--verde) 100%);
    box-shadow: 0 14px 24px -12px rgba(16,185,129,.7);
    transition: transform .2s ease, box-shadow .2s ease;
}
[data-testid="stFileUploader"] section button::after {
    content: "Seleccionar archivo";
    font-size: .92rem; font-weight: 700; letter-spacing: .01em; color: #fff;
}
[data-testid="stFileUploader"] section button * { display: none; }
[data-testid="stFileUploader"] section button:hover {
    transform: translateY(-2px);
    box-shadow: 0 18px 28px -12px rgba(59,130,246,.75);
}
[data-testid="stFileUploader"] section button:active { transform: translateY(0); }
/* archivo ya cargado */
[data-testid="stFileUploaderFile"] {
    margin-top: .6rem;
    padding: .65rem .9rem;
    background: var(--fondo-card);
    border: 1px solid var(--borde);
    border-left: 4px solid var(--verde);
    border-radius: 14px;
    color: var(--tinta);
}
[data-testid="stFileUploaderFile"] * { color: var(--tinta); }
[data-testid="stFileUploaderFile"] small { color: var(--gris); }

/* Spinner */
[data-testid="stSpinner"] * { color: var(--tinta); }

/* Botón de descarga */
.stDownloadButton > button {
    width: 100%;
    background: linear-gradient(135deg, var(--azul) 0%, var(--verde) 100%);
    color: #fff; border: none;
    padding: .95rem 1.2rem;
    border-radius: 14px;
    font-weight: 700; font-size: 1.02rem;
    box-shadow: 0 14px 26px -12px rgba(16,185,129,.7);
    transition: all .2s ease;
}
.stDownloadButton > button:hover {
    transform: translateY(-2px);
    box-shadow: 0 18px 30px -12px rgba(59,130,246,.7);
    color: #fff; border: none;
}
.stDownloadButton > button:active { transform: translateY(0); }

/* Tarjeta de resultado */
.resultado {
    background: var(--fondo-card);
    border: 1px solid var(--borde);
    border-left: 5px solid var(--verde);
    border-radius: 16px;
    padding: 1.2rem 1.3rem;
    margin: 1.2rem 0 1rem 0;
    box-shadow: 0 10px 24px -16px var(--sombra);
}
.resultado h3 { margin: 0 0 .25rem 0; font-size: 1.15rem; color: var(--tinta); }
.resultado p { margin: 0; color: var(--gris); font-size: .92rem; }
.stats { display: grid; grid-template-columns: repeat(3, 1fr); gap: .7rem; margin-top: 1rem; }
.stat {
    background: var(--fondo-suave); border: 1px solid var(--borde);
    border-radius: 12px; padding: .7rem .8rem; text-align: center;
}
.stat .valor { font-size: 1.35rem; font-weight: 800; color: var(--acento-texto); }
.stat .etq { font-size: .72rem; color: var(--gris); text-transform: uppercase; letter-spacing: .05em; }

/* Alertas nativas */
[data-testid="stAlert"] {
    border-radius: 14px;
    background: var(--fondo-card);
    border: 1px solid var(--borde);
}
[data-testid="stAlert"] * { color: var(--tinta); }

/* Pie de página */
.pie {
    text-align: center; color: var(--gris); font-size: .8rem;
    margin-top: 2.5rem; padding-top: 1.2rem; border-top: 1px solid var(--borde);
}

@media (max-width: 640px) {
    .hero { padding: 1.7rem 1.3rem; }
    .hero h1 { font-size: 2rem; }
    .hero .logo { width: 44px; height: 44px; }
    .stats { grid-template-columns: 1fr; }
}
</style>
"""

def paso(numero, titulo, subtitulo=""):
    sub = f'<div class="sub">{subtitulo}</div>' if subtitulo else ""
    st.markdown(
        f'<div class="paso"><div class="num">{numero}</div>'
        f'<div><div class="titulo">{titulo}</div>{sub}</div></div>',
        unsafe_allow_html=True,
    )


def _alternar_tema():
    st.session_state["tema_claro"] = not st.session_state.get("tema_claro", False)


@st.fragment
def selector_tema():
    # Fragmento: cambiar el tema no vuelve a procesar el archivo subido.
    claro = st.session_state.get("tema_claro", False)
    with st.container(key="selector_tema"):
        st.button(
            ":material/dark_mode:" if claro else ":material/light_mode:",
            key="btn_tema",
            on_click=_alternar_tema,
            help="Cambiar a tema oscuro" if claro else "Cambiar a tema claro",
        )
    if claro:
        st.markdown(CSS_TEMA_CLARO, unsafe_allow_html=True)


# =============================================================================
# 6. INTERFAZ WEB STREAMLIT
# =============================================================================


@st.cache_data(max_entries=2, ttl=900, show_spinner=False)
def _procesar_cacheado(datos):
    """Pulsar "Descargar" vuelve a ejecutar el script: el resultado se guarda para no reprocesar el archivo."""
    return procesar_xlsx(datos)


def main():
    st.set_page_config(page_title=f"{NOMBRE_APP} · Separador de nombres", layout="centered", page_icon=ICONO_APP)
    st.markdown(CSS_TEMA_OSCURO + ESTILOS, unsafe_allow_html=True)

    st.markdown(
        f"""
        <div class="hero">
            <div class="badge">Herramienta para Excel</div>
            <div class="marca">{LOGO_SVG}<h1>{NOMBRE_APP}</h1></div>
            <div class="lema">{LEMA_APP}</div>
            <p>Divide automáticamente los nombres completos en 4 columnas. Deduce solo, fila por fila,
            si cada nombre viene como Nombres + Apellidos o Apellidos + Nombres, y deja el documento
            tal cual: imágenes, logos, colores, tablas y filtros.</p>
            <div class="chips">
                <span class="chip">✔ Entiende (A) y (B) por celda</span>
                <span class="chip">✔ Deduce el orden sin letras</span>
                <span class="chip">✔ Conserva imágenes y formato</span>
                <span class="chip">✔ Detecta conectores (DE LA, DEL…)</span>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    paso(1, "Sube tu archivo", "Formato Excel (.xlsx) · no necesitas indicar nada más")
    archivo_subido = st.file_uploader(
        "Sube tu archivo Excel (.xlsx)", type=["xlsx"], label_visibility="collapsed"
    )

    if archivo_subido is not None:
        paso(2, "Resultado", "Revisa y descarga tu archivo procesado")
        if archivo_subido.size > LIMITE_SUBIDA_MB * 1024 * 1024:
            st.error(f"❌ El archivo pesa más de {LIMITE_SUBIDA_MB} MB. Divídelo en partes más pequeñas e inténtalo de nuevo.")
        else:
            with st.spinner("Procesando el archivo sin alterar su formato..."):
                try:
                    salida, res = _procesar_cacheado(archivo_subido.getvalue())

                    c = res["conteo"]
                    n_revisar = len(res["revisar"])
                    detalle = (f"Con (A): {c['A']} · con (B): {c['B']} · "
                               f"sin letra, deducidas por el programa: {c['AUTO']}")

                    st.markdown(
                        f"""
                        <div class="resultado">
                            <h3>✅ ¡Proceso completado!</h3>
                            <p>{detalle}</p>
                            <div class="stats">
                                <div class="stat"><div class="valor">{res['registros']}</div><div class="etq">Registros</div></div>
                                <div class="stat"><div class="valor">4</div><div class="etq">Columnas nuevas</div></div>
                                <div class="stat"><div class="valor">{n_revisar}</div><div class="etq">Por revisar</div></div>
                            </div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
                    for aviso in res["advertencias"]:
                        st.info(aviso)
                    if n_revisar:
                        st.warning(
                            f"{n_revisar} fila(s) tenían poca evidencia (palabras que no están en los "
                            "diccionarios). El programa las separó con su mejor criterio; conviene mirarlas."
                        )
                        with st.expander("Filas para revisar"):
                            st.dataframe(res["revisar"], hide_index=True)
                    if res["vista_previa"]:
                        with st.expander("Vista previa de las primeras filas"):
                            st.dataframe(res["vista_previa"], hide_index=True)
                    st.download_button(
                        label="📥 Descargar Archivo Procesado",
                        data=salida,
                        file_name=f"OK_{archivo_subido.name}",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                except (ColumnaNoEncontrada, ArchivoNoSoportado) as e:
                    st.error(f"❌ {e}")
                except Exception as e:
                    st.error(f"Error procesando el archivo: {e}")

    st.markdown(f'<div class="pie">{NOMBRE_APP} · Tus archivos se procesan en memoria y no se almacenan</div>', unsafe_allow_html=True)

    selector_tema()


if __name__ == "__main__":
    main()
