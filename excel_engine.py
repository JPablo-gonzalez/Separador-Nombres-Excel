"""Motor de procesamiento de Excel.

Inserta las 4 columnas nuevas editando el XML del .xlsx en el lugar: desplaza celdas, anchos,
combinaciones, tablas, filtros, fórmulas, formatos condicionales, validaciones, nombres definidos
y dibujos, y copia sin cambios imágenes, logos, estilos y tema. openpyxl se usa para traducir
fórmulas y coordenadas.
"""

import io
import posixpath
import re
import shutil
import zipfile

from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

import edades
from parser import resolver_lista

# =============================================================================
# 1. REFERENCIAS: desplazamiento de columnas (funciones puras sobre texto)
# =============================================================================
# Al insertar columnas, toda referencia a una columna >= ins se desplaza n columnas a la derecha,
# igual que hace Excel. Cada extremo de un rango se evalúa por separado: un rango que cruza el
# punto de inserción se ensancha.

_RE_CELDA = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_RE_COL = re.compile(r"^(\$?)([A-Za-z]{1,3})$")


def _desplazar_col(letras, ins, n):
    c = column_index_from_string(letras.upper())
    return get_column_letter(c + n) if c >= ins else letras


def _desplazar_operando(op, titulo, ins, n, hoja_propia):
    if "!" in op:
        pref, ref = op.rsplit("!", 1)
        if pref.strip("'").replace("''", "'") != titulo:
            return op                      # apunta a otra hoja
    else:
        pref, ref = None, op
        if not hoja_propia:
            return op                      # sin hoja explícita, pero la fórmula es de otra hoja
    if "[" in ref:
        return op                          # referencias estructuradas de tabla (por nombre)
    partes = ref.split(":")
    if len(partes) > 2:
        return op
    nuevas = []
    for p in partes:
        m = _RE_CELDA.match(p)
        if m:
            nuevas.append(f"{m.group(1)}{_desplazar_col(m.group(2), ins, n)}{m.group(3)}{m.group(4)}")
            continue
        m = _RE_COL.match(p)
        if m and len(partes) == 2:         # columnas completas (A:C); solas podrían ser un nombre
            nuevas.append(f"{m.group(1)}{_desplazar_col(m.group(2), ins, n)}")
            continue
        nuevas.append(p)
    ref_nueva = ":".join(nuevas)
    return f"{pref}!{ref_nueva}" if pref is not None else ref_nueva


def _desplazar_formula(texto, titulo, ins, n, hoja_propia):
    """texto incluye el '=' inicial. Si algo falla, devuelve el texto original."""
    try:
        tok = Tokenizer(texto)
        cambio = False
        for t in tok.items:
            if t.type == Token.OPERAND and t.subtype == Token.RANGE:
                nuevo = _desplazar_operando(t.value, titulo, ins, n, hoja_propia)
                if nuevo != t.value:
                    t.value = nuevo
                    cambio = True
        return tok.render() if cambio else texto
    except Exception:
        return texto


def _desplazar_rangos(texto, titulo, ins, n):
    """Lista de rangos separados por espacios (sqref) de la hoja propia."""
    return " ".join(_desplazar_operando(r, titulo, ins, n, True) for r in str(texto).split())


# =============================================================================
# 2. PROCESAMIENTO DIRECTO DEL .XLSX
# =============================================================================
# El .xlsx es un ZIP de archivos XML. En vez de cargarlo entero con openpyxl (que reconstruye el libro,
# crea una celda en memoria por cada fila declarada y pierde formas, imágenes y partes de Office),
# se edita el XML de la hoja en el lugar:
#
#   · solo se recorren las filas que EXISTEN en el archivo (una celda perdida en la fila 1.048.255
#     ya no cuesta un millón de iteraciones);
#   · imágenes, logos, tema, estilos y cualquier otra parte se copian byte a byte;
#   · solo se reescribe lo que cambia al insertar 4 columnas: celdas, anchos, combinaciones,
#     tablas, filtros, fórmulas, formatos condicionales, validaciones, nombres definidos y
#     la posición de los dibujos.

HEADERS_NUEVOS = ["PRIMER APELLIDO", "SEGUNDO APELLIDO", "PRIMER NOMBRE", "SEGUNDO NOMBRE"]
N_NUEVAS = len(HEADERS_NUEVOS)
ANCHO_NUEVAS = 19

LIMITE_SUBIDA_MB = 30            # tamaño máximo del .xlsx subido
LIMITE_DESCOMPRIMIDO_MB = 200    # tamaño máximo del contenido descomprimido (protege de "zip bombs")
FILA_MAX_ENCABEZADO = 30         # el encabezado se busca por encima de esta fila (como antes)


class ColumnaNoEncontrada(Exception):
    pass


class ArchivoNoSoportado(Exception):
    """Archivo inválido o fuera de los límites de seguridad (el mensaje se muestra al usuario)."""


# ---------------------------------------------------------------- utilidades de texto / XML

_RE_ENT_NUM = re.compile(r"&#(x[0-9A-Fa-f]+|[0-9]+);")
_RE_ESC_X = re.compile(r"_x([0-9A-Fa-f]{4})_")
_RE_ATTR = re.compile(rb"""([\w:.-]+)\s*=\s*(?:"([^"]*)"|'([^']*)')""")
_RE_REF_B = re.compile(rb"^([A-Z]+)(\d+)$")
_RE_FILA = re.compile(rb"<row\b([^>]*?)(?:/>|>(.*?)</row>)", re.S)
_RE_CEL = re.compile(rb"<c\b([^>]*?)(?:/>|>(.*?)</c>)", re.S)
_RE_F = re.compile(rb"<f\b([^>]*?)(?:/>|>(.*?)</f>)", re.S)
_RE_T = re.compile(rb"<t\b[^>]*?(?:/>|>(.*?)</t>)", re.S)
_RE_RPH = re.compile(rb"<rPh\b.*?</rPh>", re.S)
_RE_SPANS = re.compile(rb'\sspans="[^"]*"')
_RE_R_ATTR = re.compile(rb'(\s)r="[^"]*"')
_RE_R_CELDA = re.compile(rb'\sr="([A-Z]+)\d+"')
_RE_S_ATTR = re.compile(rb'\ss="(\d+)"')


def _esc(texto, atributo=False):
    texto = texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return texto.replace('"', "&quot;") if atributo else texto


def _unesc(texto):
    if "&" not in texto:
        return texto
    texto = _RE_ENT_NUM.sub(
        lambda m: chr(int(m.group(1)[1:], 16) if m.group(1)[0] in "xX" else int(m.group(1))), texto)
    return (texto.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
            .replace("&apos;", "'").replace("&amp;", "&"))


def _unesc_x(texto):
    """Excel guarda los caracteres de control como _x000D_."""
    return _RE_ESC_X.sub(lambda m: chr(int(m.group(1), 16)), texto) if "_x" in texto else texto


def _attrs(b):
    return {m.group(1): (m.group(2) if m.group(2) is not None else m.group(3)) for m in _RE_ATTR.finditer(b)}


_CACHE_IDX, _CACHE_LETRA = {}, {}


def _idx_col(letras):
    v = _CACHE_IDX.get(letras)
    if v is None:
        v = _CACHE_IDX[letras] = column_index_from_string(letras.decode("ascii"))
    return v


def _letra(col):
    v = _CACHE_LETRA.get(col)
    if v is None:
        v = _CACHE_LETRA[col] = get_column_letter(col).encode("ascii")
    return v


def _texto_t(b):
    """Texto de un fragmento con <t> (cadena compartida o inline), sin los fonéticos <rPh>."""
    if not b:
        return ""
    b = _RE_RPH.sub(b"", b)
    partes = [m.group(1) or b"" for m in _RE_T.finditer(b)]
    return _unesc_x(_unesc(b"".join(partes).decode("utf-8", "replace")))


def _leer_cadenas(datos):
    return [_texto_t(m.group(1)) for m in re.finditer(rb"<si\b[^>]*?(?:/>|>(.*?)</si>)", datos, re.S)]


def _texto_celda(cattrs, inner, cadenas):
    """Texto de una celda (None si no es texto o si es una fórmula)."""
    if not inner or b"<f" in inner:
        return None
    tipo = cattrs.get(b"t", b"n")
    if tipo == b"s":
        m = re.search(rb"<v>\s*(\d+)\s*</v>", inner)
        if m and int(m.group(1)) < len(cadenas):
            return cadenas[int(m.group(1))]
        return None
    if tipo == b"inlineStr":
        m = re.search(rb"<is\b[^>]*>(.*?)</is>", inner, re.S)
        return _texto_t(m.group(1)) if m else None
    if tipo == b"str":
        m = re.search(rb"<v>(.*?)</v>", inner, re.S)
        return _unesc_x(_unesc(m.group(1).decode("utf-8", "replace"))) if m else None
    return None


def _sub_attr(datos, nombre, fn, etiquetas=None):
    """Aplica fn(str) -> str al valor del atributo `nombre` (en cualquier etiqueta o solo en `etiquetas`)."""
    if etiquetas:
        patron = re.compile(rb"(<(?:" + b"|".join(etiquetas) + rb")\b[^>]*?\s" + nombre + rb'=")([^"]*)(")')
    else:
        patron = re.compile(rb"(\s" + nombre + rb'=")([^"]*)(")')

    def reemplazo(m):
        original = _unesc(m.group(2).decode("utf-8"))
        nuevo = fn(original)
        if nuevo == original:
            return m.group(0)
        return m.group(1) + _esc(nuevo, True).encode("utf-8") + m.group(3)

    return patron.sub(reemplazo, datos)


def _sub_texto(datos, patron, fn):
    """Patrón con tres grupos (apertura)(texto)(cierre): aplica fn(bytes) -> bytes al texto central."""
    return patron.sub(lambda m: m.group(1) + fn(m.group(2)) + m.group(3), datos)


# ---------------------------------------------------------------- plan de desplazamiento

class _Desp:
    """Qué se desplaza: n columnas a partir de la columna `ins` (base 1) de la hoja `titulo`."""

    def __init__(self, titulo, col_nombres, n=N_NUEVAS):
        self.titulo = titulo
        self.col = col_nombres
        self.ins = col_nombres + 1
        self.n = n

    def rangos(self, texto):
        return _desplazar_rangos(texto, self.titulo, self.ins, self.n)

    def formula(self, texto, propia=True):
        return _desplazar_formula("=" + texto, self.titulo, self.ins, self.n, propia)[1:]

    def formula_b(self, b, propia=True):
        """Formula en bytes (con escapes XML). Si no cambia, devuelve exactamente los mismos bytes."""
        original = _unesc(b.decode("utf-8"))
        nuevo = self.formula(original, propia)
        return b if nuevo == original else _esc(nuevo).encode("utf-8")


# ---------------------------------------------------------------- estructura del libro

def _resolver_ruta(base, destino):
    if destino.startswith("/"):
        return destino.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base), destino))


def _relaciones(zin, ruta_parte):
    """{Id: (Type, ruta_resuelta)} de las relaciones de una parte del paquete."""
    carpeta, nombre = posixpath.split(ruta_parte)
    ruta_rels = posixpath.join(carpeta, "_rels", nombre + ".rels")
    try:
        datos = zin.read(ruta_rels)
    except KeyError:
        return {}
    rels = {}
    for m in re.finditer(rb"<Relationship\b([^>]*?)/?>", datos):
        a = _attrs(m.group(1))
        if b"Id" in a and b"Target" in a and a.get(b"TargetMode") != b"External":
            rels[a[b"Id"].decode()] = (a.get(b"Type", b"").decode(), _resolver_ruta(ruta_parte, a[b"Target"].decode()))
    return rels


def _hojas(zin):
    """Hojas de cálculo del libro: lista de (nombre, ruta) y el índice de la hoja activa."""
    wb = zin.read("xl/workbook.xml")
    rels = _relaciones(zin, "xl/workbook.xml")
    hojas = []
    for m in re.finditer(rb"<sheet\b([^>]*?)/?>", wb):
        a = _attrs(m.group(1))
        rid = next((v for k, v in a.items() if k.endswith(b":id") or k == b"id"), None)
        if rid is None or rid.decode() not in rels:
            continue
        tipo, ruta = rels[rid.decode()]
        if tipo.endswith("/worksheet"):
            hojas.append((_unesc(a.get(b"name", b"").decode("utf-8")), ruta))
    if not hojas:
        raise ArchivoNoSoportado("El libro no contiene hojas de cálculo.")
    # Índice de la hoja activa entre TODAS las hojas (como openpyxl: activeTab)
    todas = [(_attrs(m.group(1))) for m in re.finditer(rb"<sheet\b([^>]*?)/?>", wb)]
    ma = re.search(rb"<workbookView\b[^>]*?\bactiveTab=\"(\d+)\"", wb)
    activa = int(ma.group(1)) if ma else 0
    ruta_activa = None
    if activa < len(todas):
        rid = next((v for k, v in todas[activa].items() if k.endswith(b":id") or k == b"id"), None)
        if rid is not None and rid.decode() in rels and rels[rid.decode()][0].endswith("/worksheet"):
            ruta_activa = rels[rid.decode()][1]
    if ruta_activa is None:
        ruta_activa = hojas[0][1]
    return hojas, ruta_activa


# ---------------------------------------------------------------- columnas (<cols>) y anchos

class _Anchos:
    """Ancho en píxeles de cada columna, para recalcular dibujos que cruzan el punto de inserción."""

    def __init__(self, definiciones, defecto_px=64):
        self.defs = definiciones          # [(min, max, ancho_en_caracteres | None, oculta)]
        self.defecto = defecto_px

    @staticmethod
    def _px(ancho):
        return int(((256 * ancho + int(128 / 7)) / 256) * 7)

    def px(self, col):
        for mn, mx, ancho, oculta in self.defs:
            if mn <= col <= mx:
                if oculta:
                    return 0
                return self._px(ancho) if ancho is not None else self.defecto
        return self.defecto


def _leer_cols(head):
    m = re.search(rb"<cols>(.*?)</cols>", head, re.S)
    cols = []
    if m:
        for c in re.finditer(rb"<col\b([^>]*?)/?>", m.group(1)):
            a = _attrs(c.group(1))
            cols.append(dict(a))
    return m, cols


def _construir_cols(cols_originales, col_nombres, n=N_NUEVAS):
    """Nuevas definiciones <col>: las de la derecha se desplazan y se agregan las columnas nuevas."""
    ins = col_nombres + 1
    resultado, estilo_nombres = [], None
    for a in cols_originales:
        mn, mx = int(a[b"min"]), int(a[b"max"])
        if mn <= col_nombres <= mx and b"style" in a:
            estilo_nombres = a[b"style"]
        if mx < ins:
            trozos = [(mn, mx)]
        elif mn >= ins:
            trozos = [(mn + n, mx + n)]
        else:                              # el rango atraviesa el punto de inserción: se parte en dos
            trozos = [(mn, ins - 1), (ins + n, mx + n)]
        for a1, b1 in trozos:
            if a1 > 16384:
                continue
            nuevo = dict(a)
            nuevo[b"min"], nuevo[b"max"] = str(a1).encode(), str(min(b1, 16384)).encode()
            resultado.append(nuevo)
    nuevas = {b"min": str(ins).encode(), b"max": str(ins + n - 1).encode(),
              b"width": str(ANCHO_NUEVAS).encode(), b"customWidth": b"1"}
    if estilo_nombres is not None:
        nuevas[b"style"] = estilo_nombres
    resultado.append(nuevas)
    resultado.sort(key=lambda a: int(a[b"min"]))
    return resultado


def _serializar_cols(cols):
    partes = []
    for a in cols:
        partes.append(b"<col " + b" ".join(k + b'="' + v + b'"' for k, v in a.items()) + b"/>")
    return b"<cols>" + b"".join(partes) + b"</cols>"


def _anchos_desde(cols, defecto_px):
    return _Anchos([(int(a[b"min"]), int(a[b"max"]),
                     float(a[b"width"]) if b"width" in a else None,
                     a.get(b"hidden") in (b"1", b"true")) for a in cols], defecto_px)


def _ancho_defecto_px(head):
    m = re.search(rb"<sheetFormatPr\b([^>]*?)/?>", head)
    if m:
        a = _attrs(m.group(1))
        if b"defaultColWidth" in a:
            return _Anchos._px(float(a[b"defaultColWidth"]))
    return 64


# ---------------------------------------------------------------- cabecera y cola de la hoja

def _ajustar_xml_hoja(b, p):
    """Referencias de todo lo que rodea a <sheetData>: selección, combinaciones, filtros, formatos
    condicionales, validaciones, hipervínculos, paneles..."""
    desp = p.rangos
    b = _sub_attr(b, b"sqref", desp)
    b = _sub_attr(b, b"activeCell", desp)

    def inicio_de_vista(texto):
        # Si el desplazamiento de la hoja empezaba justo donde se insertan las columnas, se deja ahí:
        # así las 4 columnas nuevas quedan a la vista al abrir el archivo (y no escondidas a la izquierda).
        m = _RE_CELDA.match(texto)
        if m and column_index_from_string(m.group(2).upper()) == p.ins:
            return texto
        return desp(texto)
    b = _sub_attr(b, b"topLeftCell", inicio_de_vista)
    b = _sub_attr(b, b"ref", desp, etiquetas=(b"mergeCell", b"sortState", b"sortCondition", b"hyperlink"))
    b = _sub_texto(b, re.compile(rb"(<formula[12]?>)(.*?)(</formula[12]?>)", re.S), p.formula_b)
    b = _sub_texto(b, re.compile(rb"(<xm:f>)(.*?)(</xm:f>)", re.S), p.formula_b)
    b = _sub_texto(b, re.compile(rb"(<xm:sqref>)(.*?)(</xm:sqref>)", re.S),
                   lambda t: p.rangos(_unesc(t.decode("utf-8"))).encode("utf-8"))

    # Paneles inmovilizados: si la inserción cae dentro de la zona fija, esta crece
    def panel(m):
        a = _attrs(m.group(1))
        if a.get(b"state", b"").startswith(b"frozen") and b"xSplit" in a:
            x = float(a[b"xSplit"])
            if x == int(x) and p.ins <= int(x):
                return m.group(0).replace(b'xSplit="' + a[b"xSplit"] + b'"', b'xSplit="%d"' % (int(x) + p.n))
        return m.group(0)
    b = re.sub(rb"<pane\b([^>]*?)/?>", panel, b)

    # Autofiltro de hoja (también se ensancha si terminaba justo en la columna de nombres)
    def autofiltro(m):
        a = _attrs(m.group(1))
        if b"ref" not in a:
            return m.group(0)
        c1, r1, c2, r2 = range_boundaries(a[b"ref"].decode())
        n1 = c1 + p.n if c1 >= p.ins else c1
        n2 = c2 + p.n if c2 >= p.col else c2
        nuevo_ref = f"{get_column_letter(n1)}{r1}:{get_column_letter(n2)}{r2}".encode()
        etiqueta = m.group(0)
        cuerpo = m.group(2)
        etiqueta = etiqueta.replace(b'ref="' + a[b"ref"] + b'"', b'ref="' + nuevo_ref + b'"', 1)
        if cuerpo and b"filterColumn" in cuerpo:
            def columna(mc):
                cid = int(mc.group(2))
                return mc.group(1) + (str(cid + p.n).encode() if c1 + cid >= p.ins else mc.group(2)) + mc.group(3)
            nuevo_cuerpo = re.sub(rb'(<filterColumn\b[^>]*?\bcolId=")(\d+)(")', columna, cuerpo)
            etiqueta = etiqueta.replace(cuerpo, nuevo_cuerpo, 1)
        return etiqueta
    b = re.sub(rb"<autoFilter\b([^>]*?)(?:/>|>(.*?)</autoFilter>)", autofiltro, b, flags=re.S)
    return b


def _ajustar_cabecera(head, p):
    """Cabecera de la hoja: dimensión, columnas (<cols>) y referencias."""
    head = _ajustar_xml_hoja(head, p)

    def dimension(m):
        ref = m.group(2).decode()
        try:
            c1, r1, c2, r2 = range_boundaries(ref)
        except Exception:
            return m.group(0)
        n1 = c1 + p.n if c1 >= p.ins else c1
        n2 = c2 + p.n if c2 >= p.ins else c2
        if c2 >= p.col:
            n2 = max(n2, p.ins + p.n - 1)
        return m.group(1) + f"{get_column_letter(n1)}{r1}:{get_column_letter(n2)}{r2}".encode() + m.group(3)
    head = re.sub(rb'(<dimension\b[^>]*?\sref=")([^"]*)(")', dimension, head)

    m, cols = _leer_cols(head)
    defecto = _ancho_defecto_px(head)
    anchos_viejos = _anchos_desde(cols, defecto)
    nuevas = _construir_cols(cols, p.col)
    anchos_nuevos = _anchos_desde(nuevas, defecto)
    bloque = _serializar_cols(nuevas)
    if m:
        head = head[:m.start()] + bloque + head[m.end():]
    else:
        head = head + bloque              # <cols> va justo antes de <sheetData>
    return head, anchos_viejos, anchos_nuevos


# ---------------------------------------------------------------- tablas de Excel

def _ajustar_tabla(xml, p):
    """Ajusta una tabla (xl/tables/tableN.xml). Devuelve (xml, info) con info = None si no cambia el
    encabezado, o (fila_encabezado, [4 nombres]) si la tabla contiene la columna de nombres."""
    mt = re.search(rb"<table\b([^>]*)>", xml)
    if not mt:
        return xml, None
    ta = _attrs(mt.group(1))
    if b"ref" not in ta:
        return xml, None
    c1, r1, c2, r2 = range_boundaries(ta[b"ref"].decode())
    if c2 < p.col:
        return xml, None                    # tabla totalmente a la izquierda: no cambia

    filas_enc = int(ta.get(b"headerRowCount", b"1"))
    filas_tot = int(ta.get(b"totalsRowCount", b"0"))
    contiene = c1 <= p.col
    n1, n2 = (c1, c2 + p.n) if contiene else (c1 + p.n, c2 + p.n)
    L = get_column_letter
    fila_datos_max = r2 - filas_tot
    ref_tabla = f"{L(n1)}{r1}:{L(n2)}{r2}"
    ref_filtro = f"{L(n1)}{r1}:{L(n2)}{fila_datos_max}"
    ref_orden = f"{L(n1)}{r1 + filas_enc}:{L(n2)}{fila_datos_max}"

    xml = xml.replace(b'ref="' + ta[b"ref"] + b'"', b'ref="' + ref_tabla.encode() + b'"', 1)

    def filtro(m):
        etiqueta, cuerpo = m.group(0), m.group(2)
        a = _attrs(m.group(1))
        if b"ref" in a:
            etiqueta = etiqueta.replace(b'ref="' + a[b"ref"] + b'"', b'ref="' + ref_filtro.encode() + b'"', 1)
        if contiene and cuerpo and b"filterColumn" in cuerpo:
            limite = p.col - c1 + 1

            def columna(mc):
                cid = int(mc.group(2))
                return mc.group(1) + (str(cid + p.n).encode() if cid >= limite else mc.group(2)) + mc.group(3)
            etiqueta = etiqueta.replace(
                cuerpo, re.sub(rb'(<filterColumn\b[^>]*?\bcolId=")(\d+)(")', columna, cuerpo), 1)
        return etiqueta
    xml = re.sub(rb"<autoFilter\b([^>]*?)(?:/>|>(.*?)</autoFilter>)", filtro, xml, count=1, flags=re.S)

    def orden(m):
        a = _attrs(m.group(1))
        etiqueta = m.group(0)
        if b"ref" in a:
            etiqueta = etiqueta.replace(b'ref="' + a[b"ref"] + b'"', b'ref="' + ref_orden.encode() + b'"', 1)
        return etiqueta
    xml = re.sub(rb"<sortState\b([^>]*?)(?:/>|>.*?</sortState>)",
                 lambda m: _sub_attr(orden(m), b"ref", p.rangos, etiquetas=(b"sortCondition",)), xml, flags=re.S)

    nuevo_info = None
    mc = re.search(rb"<tableColumns\b([^>]*)>(.*?)</tableColumns>", xml, re.S)
    if mc:
        columnas = [m.group(0) for m in re.finditer(rb"<tableColumn\b[^>]*?(?:/>|>.*?</tableColumn>)", mc.group(2), re.S)]
        columnas = [_sub_texto(c, re.compile(rb"(<calculatedColumnFormula\b[^>]*>)(.*?)(</calculatedColumnFormula>)", re.S),
                               p.formula_b) for c in columnas]
        if contiene:
            usados = {_unesc(_attrs(c)[b"name"].decode("utf-8")).lower() for c in columnas if b"name" in _attrs(c)}
            ids = [int(_attrs(c).get(b"id", b"0")) for c in columnas]
            siguiente = max(ids + [0]) + 1
            nombres = []
            for h in HEADERS_NUEVOS:
                nombre, k = h, 1
                while nombre.lower() in usados:      # Excel exige nombres únicos, sin distinguir mayúsculas
                    nombre = f"{h}_{k}"
                    k += 1
                usados.add(nombre.lower())
                nombres.append(nombre)
            nuevas = [b'<tableColumn id="%d" name="%b"/>' % (siguiente + i, _esc(nm, True).encode("utf-8"))
                      for i, nm in enumerate(nombres)]
            pos = p.col - c1 + 1
            columnas = columnas[:pos] + nuevas + columnas[pos:]
            if filas_enc > 0:
                nuevo_info = (r1, nombres)
        abre = re.sub(rb'\bcount="\d+"', b'count="%d"' % len(columnas), b"<tableColumns" + mc.group(1) + b">")
        xml = xml[:mc.start()] + abre + b"".join(columnas) + b"</tableColumns>" + xml[mc.end():]
    return xml, nuevo_info


# ---------------------------------------------------------------- dibujos, comentarios, gráficos

def _ajustar_dibujo(xml, p, anchos_viejos, anchos_nuevos):
    """Mueve los objetos (imágenes, logos, formas, gráficos) que quedan a la derecha del punto de
    inserción. Los que lo cruzan se estiran o conservan su tamaño según su propiedad editAs."""
    i0 = p.ins - 1                                   # índice base 0 de la primera columna desplazada
    re_col = re.compile(rb"(<(?:\w+:)?col>)(\d+)(</(?:\w+:)?col>)")
    re_off = re.compile(rb"(<(?:\w+:)?colOff>)(-?\d+)(</(?:\w+:)?colOff>)")
    re_punto = re.compile(rb"(<(?:\w+:)?(from|to)>)(.*?)(</(?:\w+:)?\2>)", re.S)

    def leer(cuerpo):
        c = int(re_col.search(cuerpo).group(2))
        mo = re_off.search(cuerpo)
        return c, int(mo.group(2)) if mo else 0

    def escribir(cuerpo, col, off):
        cuerpo = re_col.sub(lambda m: m.group(1) + str(col).encode() + m.group(3), cuerpo, count=1)
        return re_off.sub(lambda m: m.group(1) + str(off).encode() + m.group(3), cuerpo, count=1)

    def ancla(m):
        prefijo, tipo, attrs, cuerpo = m.group(1), m.group(2), m.group(3), m.group(4)
        puntos = {mm.group(2): mm for mm in re_punto.finditer(cuerpo)}
        if b"from" not in puntos:
            return m.group(0)
        f_col, f_off = leer(puntos[b"from"].group(3))
        cambios = {}                                   # b"from" / b"to" -> nuevo contenido
        if tipo == b"oneCellAnchor" or b"to" not in puntos:
            if f_col >= i0:
                cambios[b"from"] = escribir(puntos[b"from"].group(3), f_col + p.n, f_off)
        else:
            t_col, t_off = leer(puntos[b"to"].group(3))
            edita = _attrs(attrs).get(b"editAs", b"twoCell")
            if f_col >= i0:                           # entero a la derecha: se mueve
                cambios[b"from"] = escribir(puntos[b"from"].group(3), f_col + p.n, f_off)
                cambios[b"to"] = escribir(puntos[b"to"].group(3), t_col + p.n, t_off)
            elif t_col >= i0:                         # cruza el punto de inserción
                if edita == b"twoCell":               # mover y cambiar de tamaño: se estira
                    cambios[b"to"] = escribir(puntos[b"to"].group(3), t_col + p.n, t_off)
                else:                                 # conserva su tamaño en píxeles
                    emu = sum(anchos_viejos.px(c + 1) * 9525 for c in range(f_col, t_col)) - f_off + t_off
                    col, pos, resto = f_col, f_off, emu
                    while True:
                        ancho = anchos_nuevos.px(col + 1) * 9525
                        if ancho <= 0 or pos + resto <= ancho or col > 16383:
                            break
                        resto -= ancho - pos
                        pos, col = 0, col + 1
                    cambios[b"to"] = escribir(puntos[b"to"].group(3), col, pos + resto)
        if not cambios:
            return m.group(0)
        nuevo = cuerpo
        for clave in sorted(cambios, key=lambda k: -puntos[k].start(3)):      # de atrás hacia adelante
            mm = puntos[clave]
            nuevo = nuevo[:mm.start(3)] + cambios[clave] + nuevo[mm.end(3):]
        return b"<" + prefijo + tipo + attrs + b">" + nuevo + b"</" + prefijo + tipo + b">"

    return re.sub(rb"<((?:\w+:)?)(twoCellAnchor|oneCellAnchor)\b([^>]*)>(.*?)</\1\2>", ancla, xml, flags=re.S)


def _ajustar_vml(xml, p):
    i0 = p.ins - 1

    def columna(m):
        v = int(m.group(2))
        return m.group(1) + str(v + p.n if v >= i0 else v).encode() + m.group(3)

    def ancla(m):
        partes = [x.strip() for x in m.group(2).split(b",")]
        for i in (0, 4):
            if len(partes) > i and partes[i].isdigit() and int(partes[i]) >= i0:
                partes[i] = str(int(partes[i]) + p.n).encode()
        return m.group(1) + b", ".join(partes) + m.group(3)

    xml = re.sub(rb"(<x:Column>)(\d+)(</x:Column>)", columna, xml)
    return re.sub(rb"(<x:Anchor>)(.*?)(</x:Anchor>)", ancla, xml, flags=re.S)


# ---------------------------------------------------------------- celdas nuevas y reescritura de filas

def _celda_nueva(col, fila, texto, estilo):
    ref = _letra(col) + str(fila).encode()
    s = b' s="' + estilo + b'"' if estilo and estilo != b"0" else b""
    if not texto:
        return b'<c r="' + ref + b'"' + s + b"/>"
    conservar = b' xml:space="preserve"' if (texto != texto.strip() or "  " in texto) else b""
    return (b'<c r="' + ref + b'"' + s + b' t="inlineStr"><is><t' + conservar + b">"
            + _esc(texto).encode("utf-8") + b"</t></is></c>")


class _Plan(_Desp):
    """Todo lo necesario para reescribir las filas de la hoja."""

    def __init__(self, titulo, col_nombres, fila_header):
        super().__init__(titulo, col_nombres)
        self.fila_header = fila_header
        self.cabeceras = {fila_header: list(HEADERS_NUEVOS)}   # fila -> 4 títulos
        self.resultados = {}                                   # fila -> Resultado
        self.expandir = set()                                  # grupos de fórmulas compartidas a expandir
        self.maestras = {}                                     # si -> (coordenada, fórmula)


def _formula_celda(inner, fila, col, p):
    """Reescribe el <f> de una celda: desplaza sus referencias y expande las fórmulas compartidas que
    quedarían partidas por la inserción (las de varias columnas)."""
    def f(m):
        a = _attrs(m.group(1))
        texto = m.group(2)
        if a.get(b"t") == b"shared" and a.get(b"si") in p.expandir:
            if texto is not None:                                   # la maestra pasa a fórmula normal
                return b"<f>" + p.formula_b(texto) + b"</f>"
            coord, base = p.maestras[a[b"si"]]                      # un hijo: se traduce desde la maestra
            nueva = Translator("=" + base, origin=coord).translate_formula(f"{get_column_letter(col)}{fila}")[1:]
            return b"<f>" + _esc(p.formula(nueva)).encode("utf-8") + b"</f>"
        if texto is None:
            return m.group(0)                                       # hijo de un grupo compartido: no cambia
        atributos = m.group(1)
        if b"ref" in a:                                             # maestra compartida o fórmula matricial
            nuevo_ref = p.rangos(a[b"ref"].decode()).encode()
            atributos = atributos.replace(b'ref="' + a[b"ref"] + b'"', b'ref="' + nuevo_ref + b'"', 1)
        return b"<f" + atributos + b">" + p.formula_b(texto) + b"</f>"
    return _RE_F.sub(f, inner)


def _reescribir_fila(fila_attrs, inner, fila, p):
    """Desplaza las celdas de la derecha, inserta las 4 nuevas junto a la de nombres y devuelve la fila."""
    celdas = list(_RE_CEL.finditer(inner))
    resto = _RE_CEL.sub(b"", inner).strip()           # contenido que no son celdas (extLst...), se conserva
    partes, insertadas, estilo_nombre, hay_nombre = [], False, None, False
    siguiente = 1
    ins, n = p.ins, p.n
    for cm in celdas:
        cattrs, cinner = cm.group(1), cm.group(2)
        mr = _RE_R_CELDA.search(cattrs)
        col = _idx_col(mr.group(1)) if mr else siguiente
        siguiente = col + 1
        if col == p.col:
            hay_nombre = True
            ms = _RE_S_ATTR.search(cattrs)
            estilo_nombre = ms.group(1) if ms else None
        if col >= ins and not insertadas:
            if hay_nombre:
                partes.append(_celdas_para(fila, estilo_nombre, p))
            insertadas = True
        if col >= ins:
            cattrs = _RE_R_ATTR.sub(lambda m: m.group(1) + b'r="' + _letra(col + n) + str(fila).encode() + b'"', cattrs, count=1)
        if cinner is not None and b"<f" in cinner:
            cinner = _formula_celda(cinner, fila, col, p)
        if cinner is None:
            partes.append(b"<c" + cattrs + b"/>")
        else:
            partes.append(b"<c" + cattrs + b">" + cinner + b"</c>")
    if hay_nombre and not insertadas:
        partes.append(_celdas_para(fila, estilo_nombre, p))
    fila_attrs = _RE_SPANS.sub(b"", fila_attrs)         # 'spans' es solo una pista de optimización
    return b"<row" + fila_attrs + b">" + b"".join(partes) + resto + b"</row>"


def _celdas_para(fila, estilo, p):
    if fila in p.cabeceras:
        textos = p.cabeceras[fila]
    elif fila in p.resultados:
        r = p.resultados[fila]
        textos = [r.p_ape, r.s_ape, r.p_nom, r.s_nom]
    else:
        textos = None
    if textos is None and (not estilo or estilo == b"0"):
        return b""                                       # nada que escribir ni formato que continuar
    return b"".join(_celda_nueva(p.ins + i, fila, textos[i] if textos else "", estilo) for i in range(p.n))


# ---------------------------------------------------------------- recorrido de la hoja

def _numero_fila(fattrs, anterior):
    r = fattrs.get(b"r")
    return int(r) if r else anterior + 1


def _buscar_encabezado(hoja, ini, fin, cadenas):
    """Primera celda con 'NOMBRE' o 'APELLIDO' por encima de FILA_MAX_ENCABEZADO (lectura por filas,
    de izquierda a derecha, igual que la versión anterior). Devuelve (fila, columna, textos_de_la_fila)."""
    anterior = 0
    for fm in _RE_FILA.finditer(hoja, ini, fin):
        fila = _numero_fila(_attrs(fm.group(1)), anterior)
        anterior = fila
        if fila >= FILA_MAX_ENCABEZADO:
            break
        if fm.group(2) is None:
            continue
        textos, hallada, siguiente = [], None, 1
        for cm in _RE_CEL.finditer(fm.group(2)):
            ca = _attrs(cm.group(1))
            mr = _RE_REF_B.match(ca.get(b"r", b""))
            col = _idx_col(mr.group(1)) if mr else siguiente
            siguiente = col + 1
            val = (_texto_celda(ca, cm.group(2), cadenas) or "").upper()
            textos.append(val)
            if hallada is None and ("NOMBRE" in val or "APELLIDO" in val) and val not in HEADERS_NUEVOS:
                hallada = col
        if hallada is not None:
            return fila, hallada, textos
    return None


def _recoger_nombres(hoja, ini, fin, cadenas, p, filas_omitir):
    """Textos de la columna de nombres debajo del encabezado y grupos de fórmulas compartidas."""
    nombres, anterior = [], 0
    hay_compartidas = hoja.find(b't="shared"', ini, fin) != -1
    multicolumna = set()
    for fm in _RE_FILA.finditer(hoja, ini, fin):
        fila = _numero_fila(_attrs(fm.group(1)), anterior)
        anterior = fila
        if fm.group(2) is None:
            continue
        inner = fm.group(2)
        buscar_nombre = fila > p.fila_header and fila not in filas_omitir
        if not buscar_nombre and not (hay_compartidas and b't="shared"' in inner):
            continue
        siguiente = 1
        for cm in _RE_CEL.finditer(inner):
            ca = _attrs(cm.group(1))
            mr = _RE_REF_B.match(ca.get(b"r", b""))
            col = _idx_col(mr.group(1)) if mr else siguiente
            siguiente = col + 1
            cinner = cm.group(2)
            if buscar_nombre and col == p.col:
                texto = _texto_celda(ca, cinner, cadenas)
                if texto is not None and texto.strip():
                    nombres.append((fila, texto))
            if hay_compartidas and cinner and b't="shared"' in cinner:
                for fmm in _RE_F.finditer(cinner):
                    fa = _attrs(fmm.group(1))
                    if fa.get(b"t") == b"shared" and b"ref" in fa and fmm.group(2) is not None:
                        c1, _, c2, _ = range_boundaries(fa[b"ref"].decode())
                        si = fa[b"si"]
                        p.maestras[si] = (f"{get_column_letter(col)}{fila}", _unesc(fmm.group(2).decode("utf-8")))
                        if c1 != c2:
                            multicolumna.add(si)
            if col > p.col and not hay_compartidas:
                break
    p.expandir = multicolumna
    return nombres


def _filas_de_totales(tablas_xml):
    omitir = set()
    for xml in tablas_xml:
        mt = re.search(rb"<table\b([^>]*)>", xml)
        if not mt:
            continue
        ta = _attrs(mt.group(1))
        tot = int(ta.get(b"totalsRowCount", b"0"))
        if tot and b"ref" in ta:
            _, _, _, r2 = range_boundaries(ta[b"ref"].decode())
            omitir.update(range(r2 - tot + 1, r2 + 1))
    return omitir


def _escribir_filas(hoja, ini, fin, p, destino):
    """Reescribe las filas existentes y las escribe por tandas (memoria acotada)."""
    lote, tam, anterior = [], 0, 0
    for fm in _RE_FILA.finditer(hoja, ini, fin):
        fila = _numero_fila(_attrs(fm.group(1)), anterior)
        anterior = fila
        if fm.group(2) is None:
            fragmento = b"<row" + _RE_SPANS.sub(b"", fm.group(1)) + b"/>"
        else:
            fragmento = _reescribir_fila(fm.group(1), fm.group(2), fila, p)
        lote.append(fragmento)
        tam += len(fragmento)
        if tam > (1 << 20):
            destino.write(b"".join(lote))
            lote, tam = [], 0
    if lote:
        destino.write(b"".join(lote))


# ---------------------------------------------------------------- fecha de nacimiento (solo lectura)
# Las fechas se leen solo para el resumen de edades en pantalla: el archivo de salida no cambia.

def _buscar_fecha(hoja, ini, fin, cadenas, fila_header):
    """Columna de FECHA DE NACIMIENTO (o variantes) por encima de FILA_MAX_ENCABEZADO. Gana el encabezado
    más claro; a igual claridad, el de la fila de encabezado de los nombres. Devuelve (fila, columna) o None."""
    mejor, anterior = None, 0
    for fm in _RE_FILA.finditer(hoja, ini, fin):
        fila = _numero_fila(_attrs(fm.group(1)), anterior)
        anterior = fila
        if fila >= FILA_MAX_ENCABEZADO:
            break
        if fm.group(2) is None:
            continue
        siguiente = 1
        for cm in _RE_CEL.finditer(fm.group(2)):
            ca = _attrs(cm.group(1))
            mr = _RE_REF_B.match(ca.get(b"r", b""))
            col = _idx_col(mr.group(1)) if mr else siguiente
            siguiente = col + 1
            texto = _texto_celda(ca, cm.group(2), cadenas)
            puntos = edades.puntaje_encabezado_fecha(texto) if texto else 0
            clave = (puntos, fila == fila_header)
            if puntos and (mejor is None or clave > mejor[0]):
                mejor = (clave, fila, col)
    return (mejor[1], mejor[2]) if mejor else None


def _valor_fecha(cattrs, inner, cadenas):
    """Valor de una celda de fecha: número (fecha de Excel), texto, o None si está vacía."""
    if not inner:
        return None
    tipo = cattrs.get(b"t", b"n")
    if tipo in (b"s", b"inlineStr") and b"<f" not in inner:
        return _texto_celda(cattrs, inner, cadenas)
    m = re.search(rb"<v>(.*?)</v>", inner, re.S)       # valor (o resultado guardado de una fórmula)
    if not m:
        return None
    texto = _unesc_x(_unesc(m.group(1).decode("utf-8", "replace")))
    if tipo == b"n":
        try:
            return float(texto)
        except ValueError:
            return texto
    return texto if tipo in (b"str", b"d") else "?"     # booleano o error: hay algo, pero no es una fecha


def _leer_fechas(zin, hoja, ini, fin, cadenas, fila_header, filas):
    """Valor de la celda de fecha de nacimiento de cada fila de `filas` (None si está vacía), o None si
    la hoja no tiene esa columna. También indica si el libro usa el sistema de fechas de 1904."""
    hallada = _buscar_fecha(hoja, ini, fin, cadenas, fila_header)
    if hallada is None:
        return None
    _, col_fecha = hallada
    buscadas, valores, anterior = set(filas), {}, 0
    for fm in _RE_FILA.finditer(hoja, ini, fin):
        fila = _numero_fila(_attrs(fm.group(1)), anterior)
        anterior = fila
        if fila not in buscadas or fm.group(2) is None:
            continue
        siguiente = 1
        for cm in _RE_CEL.finditer(fm.group(2)):
            ca = _attrs(cm.group(1))
            mr = _RE_REF_B.match(ca.get(b"r", b""))
            col = _idx_col(mr.group(1)) if mr else siguiente
            siguiente = col + 1
            if col == col_fecha:
                valores[fila] = _valor_fecha(ca, cm.group(2), cadenas)
                break
            if col > col_fecha:
                break
    fecha1904 = re.search(rb"<workbookPr\b[^>]*\bdate1904=\"(1|true)\"", zin.read("xl/workbook.xml")) is not None
    return {"valores": [valores.get(f) for f in filas], "fecha1904": fecha1904}


# ---------------------------------------------------------------- procesamiento del paquete

def _copiar_info(info):
    nueva = zipfile.ZipInfo(info.filename, date_time=info.date_time)
    nueva.compress_type = info.compress_type
    nueva.external_attr = info.external_attr
    nueva.create_system = info.create_system
    return nueva


def procesar_xlsx(datos):
    """Separa los nombres de un .xlsx editando su XML. Devuelve (bytes_del_nuevo_xlsx, estadísticas).

    No carga el libro en memoria: solo recorre las filas que existen y copia byte a byte todo lo que
    no cambia (imágenes, logos, tema, estilos...).

    Todo ocurre en memoria (io.BytesIO): no se crea ningún archivo en disco, ni siquiera temporal."""
    try:
        zin = zipfile.ZipFile(io.BytesIO(datos))
    except zipfile.BadZipFile:
        raise ArchivoNoSoportado("El archivo no es un .xlsx válido (o está dañado).")
    # El lector se cierra al terminar o al fallar: no queda ningún objeto abierto con el archivo subido.
    with zin:
        return _procesar_paquete(zin)


def _procesar_paquete(zin):
    infos = zin.infolist()
    if sum(i.file_size for i in infos) > LIMITE_DESCOMPRIMIDO_MB * 1024 * 1024:
        raise ArchivoNoSoportado(
            f"El contenido del archivo supera {LIMITE_DESCOMPRIMIDO_MB} MB descomprimido; es demasiado pesado para procesarlo aquí.")
    existentes = {i.filename for i in infos}
    if "xl/workbook.xml" not in existentes:
        raise ArchivoNoSoportado("El archivo no parece un libro de Excel (.xlsx).")

    hojas, ruta_hoja = _hojas(zin)
    titulo = next(nombre for nombre, ruta in hojas if ruta == ruta_hoja)
    hoja = zin.read(ruta_hoja)
    cadenas = _leer_cadenas(zin.read("xl/sharedStrings.xml")) if "xl/sharedStrings.xml" in existentes else []

    ini_tag = hoja.find(b"<sheetData")
    if ini_tag == -1:
        raise ArchivoNoSoportado("La hoja tiene un formato interno que no se puede procesar. Guárdala de nuevo desde Excel.")
    fin_tag = hoja.find(b">", ini_tag)
    fin_datos = hoja.rfind(b"</sheetData>")
    if hoja[fin_tag - 1:fin_tag] == b"/" or fin_datos == -1:
        raise ColumnaNoEncontrada("La hoja no tiene datos.")
    ini_datos = fin_tag + 1

    encabezado = _buscar_encabezado(hoja, ini_datos, fin_datos, cadenas)
    if encabezado is None:
        raise ColumnaNoEncontrada("No se encontró una columna válida de Nombres.")
    fila_header, col_nombres, textos_enc = encabezado
    p = _Plan(titulo, col_nombres, fila_header)

    # Partes relacionadas con la hoja
    rels = _relaciones(zin, ruta_hoja)
    ruta_tablas = [r for t, r in rels.values() if t.endswith("/table") and r in existentes]
    tablas = {r: zin.read(r) for r in ruta_tablas}
    filas_omitir = _filas_de_totales(tablas.values())

    nombres = _recoger_nombres(hoja, ini_datos, fin_datos, cadenas, p, filas_omitir)
    resultados = resolver_lista([t for _, t in nombres])
    p.resultados = {fila: res for (fila, _), res in zip(nombres, resultados)}
    fechas = _leer_fechas(zin, hoja, ini_datos, fin_datos, cadenas, fila_header, [f for f, _ in nombres])

    nuevas_partes = {}
    for ruta, xml in tablas.items():
        nuevo, info = _ajustar_tabla(xml, p)
        nuevas_partes[ruta] = nuevo
        if info:
            p.cabeceras[info[0]] = info[1]

    # Cabecera y cola de la hoja
    cabecera, anchos_viejos, anchos_nuevos = _ajustar_cabecera(hoja[:ini_tag], p)
    cola = _ajustar_xml_hoja(hoja[fin_datos + len(b"</sheetData>"):], p)

    # Dibujos, comentarios y gráficos
    for tipo, ruta in rels.values():
        if ruta not in existentes:
            continue
        if tipo.endswith("/drawing"):
            nuevas_partes[ruta] = _ajustar_dibujo(zin.read(ruta), p, anchos_viejos, anchos_nuevos)
        elif tipo.endswith("/comments") or tipo.endswith("/threadedComment"):
            nuevas_partes[ruta] = _sub_attr(zin.read(ruta), b"ref", p.rangos, etiquetas=(b"comment", b"threadedComment"))
        elif tipo.endswith("/vmlDrawing"):
            nuevas_partes[ruta] = _ajustar_vml(zin.read(ruta), p)

    # Nombres definidos y otras hojas que apuntan a esta
    wb = zin.read("xl/workbook.xml")
    nuevo_wb = _sub_texto(wb, re.compile(rb"(<definedName\b[^>]*>)(.*?)(</definedName>)", re.S),
                          lambda t: p.formula_b(t, propia=False))
    if nuevo_wb != wb:
        nuevas_partes["xl/workbook.xml"] = nuevo_wb
    marca = titulo.encode("utf-8")
    for nombre, ruta in hojas:
        if ruta == ruta_hoja or ruta not in existentes:
            continue
        otra = zin.read(ruta)
        if marca in otra and b"<f" in otra:
            nueva = _sub_texto(otra, re.compile(rb"(<f\b[^>]*>)(.*?)(</f>)", re.S), lambda t: p.formula_b(t, propia=False))
            if nueva != otra:
                nuevas_partes[ruta] = nueva
    for nombre in existentes:
        if re.match(r"xl/charts/chart\d+\.xml$", nombre):
            g = zin.read(nombre)
            if marca in g:
                nuevo = _sub_texto(g, re.compile(rb"(<c:f>)(.*?)(</c:f>)", re.S), lambda t: p.formula_b(t, propia=False))
                if nuevo != g:
                    nuevas_partes[nombre] = nuevo
        elif re.match(r"xl/pivotCache/pivotCacheDefinition\d+\.xml$", nombre):
            g = zin.read(nombre)
            nuevo = re.sub(rb"<worksheetSource\b[^>]*?/?>",
                           lambda m: (_sub_attr(m.group(0), b"ref", p.rangos)
                                      if _unesc(_attrs(m.group(0)).get(b"sheet", b"").decode("utf-8")) == titulo else m.group(0)), g)
            if nuevo != g:
                nuevas_partes[nombre] = nuevo

    # La cadena de cálculo quedaría desactualizada: se elimina y Excel la reconstruye al abrir
    quitar = set()
    if "xl/calcChain.xml" in existentes:
        quitar.add("xl/calcChain.xml")
        ct = zin.read("[Content_Types].xml")
        nuevas_partes["[Content_Types].xml"] = re.sub(rb'<Override\b[^>]*calcChain[^>]*?/>', b"", ct)
        rw = zin.read("xl/_rels/workbook.xml.rels")
        nuevas_partes["xl/_rels/workbook.xml.rels"] = re.sub(rb'<Relationship\b[^>]*calcChain[^>]*?/>', b"", rw)

    # Escritura del nuevo paquete: lo que no cambia se copia sin tocar
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in infos:
            if info.filename in quitar:
                continue
            nueva_info = _copiar_info(info)
            if info.filename == ruta_hoja:
                with zout.open(nueva_info, "w") as destino:
                    destino.write(cabecera + hoja[ini_tag:ini_datos])
                    _escribir_filas(hoja, ini_datos, fin_datos, p, destino)
                    destino.write(b"</sheetData>" + cola)
            elif info.filename in nuevas_partes:
                zout.writestr(nueva_info, nuevas_partes[info.filename])
            elif info.is_dir():
                zout.writestr(nueva_info, b"")
            else:
                with zin.open(info) as origen, zout.open(nueva_info, "w") as destino:
                    shutil.copyfileobj(origen, destino, 1 << 20)

    # Estadísticas y avisos
    conteo = {"A": 0, "B": 0, "AUTO": 0}
    vista_previa, por_revisar = [], []
    for (fila, texto), res in zip(nombres, resultados):
        conteo[res.origen] += 1
        resumen = {"Fila": fila, "Original": texto,
                   "PRIMER APELLIDO": res.p_ape, "SEGUNDO APELLIDO": res.s_ape,
                   "PRIMER NOMBRE": res.p_nom, "SEGUNDO NOMBRE": res.s_nom}
        if len(vista_previa) < 10:
            vista_previa.append(resumen)
        if res.revisar:
            por_revisar.append({**resumen, "Criterio": res.motivo})

    advertencias = []
    if any(h in textos_enc for h in HEADERS_NUEVOS):
        advertencias.append(
            "La hoja ya tenía columnas con los títulos PRIMER APELLIDO / SEGUNDO APELLIDO / PRIMER NOMBRE / "
            "SEGUNDO NOMBRE. No se tocaron: las nuevas se agregaron junto a la columna de nombres, así que "
            "esos títulos aparecen dos veces.")
    sin_ajustar = sorted({etiqueta for prefijo, etiqueta in (
        ("xl/pivotTables/", "tablas dinámicas"), ("xl/slicers/", "segmentaciones"),
        ("xl/timelines/", "escalas de tiempo"), ("xl/queryTables/", "consultas de datos externos"),
        ("xl/ctrlProps/", "controles de formulario")) if any(n.startswith(prefijo) for n in existentes)})
    if sin_ajustar:
        advertencias.append(
            "El archivo contiene " + ", ".join(sin_ajustar) + ". Se conservaron, pero sus referencias a celdas "
            "no se ajustaron; si alguna queda a la derecha de la columna de nombres, revísala.")

    resultado = salida.getvalue()
    salida.close()                  # libera el búfer de trabajo: solo queda la copia que se devuelve
    return resultado, {"registros": len(nombres), "conteo": conteo, "vista_previa": vista_previa,
                       "revisar": por_revisar, "advertencias": advertencias, "fechas_nacimiento": fechas}
