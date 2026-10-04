"""Genera archivos Excel de prueba con datos ficticios para comparar la versión actual con la de stlite.

Uso: python pruebas/generar_excels.py <carpeta_destino>
Los nombres se arman al azar (semilla fija) con nombres y apellidos comunes; no son personas reales.
"""

import io
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image as PILImage, ImageDraw

NOMBRES = ["JUAN", "MARIA", "JOSE", "ANA", "LUIS", "CARMEN", "PEDRO", "LAURA", "CARLOS", "SOFIA", "DIEGO",
           "VALENTINA", "ANDRES", "CAMILA", "FELIPE", "ISABELLA", "MATEO", "DANIELA", "SANTIAGO", "PAULA",
           "MARIA JOSE", "JUAN PABLO", "MARIA DEL CARMEN", "JOSE DE JESUS", "XIOMARA", "YERALDIN", "BRAYAN"]
APELLIDOS = ["GONZALEZ", "RODRIGUEZ", "MARTINEZ", "LOPEZ", "GARCIA", "PEREZ", "SANCHEZ", "RAMIREZ", "TORRES",
             "FLORES", "RIVERA", "GOMEZ", "DIAZ", "CRUZ", "MORALES", "ORTIZ", "GUTIERREZ", "CHAVEZ", "RAMOS",
             "DE LA CRUZ", "DEL RIO", "SAN MARTIN", "VILLAMIZAR", "QUINTERO", "BUITRAGO", "ZAPATA", "OSPINA"]


def nombre_al_azar(rnd, i):
    nom = rnd.choice(NOMBRES) + ("" if rnd.random() < .35 else " " + rnd.choice(NOMBRES))
    ape = rnd.choice(APELLIDOS) + ("" if rnd.random() < .15 else " " + rnd.choice(APELLIDOS))
    forma = i % 5
    if forma == 0:
        return f"(A) {ape} {nom}"
    if forma == 1:
        return f"(B) {nom} {ape}"
    if forma == 2:
        return f"{ape} {nom}"
    return f"{nom} {ape}"


def _logo_png():
    img = PILImage.new("RGB", (120, 60), (29, 78, 216))
    ImageDraw.Draw(img).rectangle([10, 10, 110, 50], fill=(16, 185, 129))
    b = io.BytesIO()
    img.save(b, format="PNG")
    b.seek(0)
    return b


def libro_completo(n, destino):
    """Logo, título combinado, estilos, fórmulas, formato condicional, validación, comentario, segunda hoja con
    fórmulas que apuntan a la primera, nombre definido, gráfico y metadatos."""
    rnd = random.Random(n)
    wb = Workbook()
    ws = wb.active
    ws.title = "Estudiantes"
    wb.properties.title = "Listado ficticio de estudiantes"
    wb.properties.creator = "Prueba stlite"
    wb.properties.description = "Datos inventados"
    wb.properties.keywords = "prueba, paridad"

    ws.add_image(Image(_logo_png()), "A1")
    ws.merge_cells("C1:G2")
    ws["C1"] = "COLEGIO DE PRUEBA · LISTADO 2026"
    ws["C1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["C1"].fill = PatternFill("solid", fgColor="1D4ED8")
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")

    fila_h = 4
    enc = ["N°", "NOMBRE COMPLETO", "CURSO", "NOTA 1", "NOTA 2", "PROMEDIO", "ESTADO"]
    borde = Border(*(Side(style="thin", color="94A3B8"),) * 4)
    for c, t in enumerate(enc, 1):
        cel = ws.cell(fila_h, c, t)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="059669")
        cel.border = borde
        cel.alignment = Alignment(horizontal="center")
    for i in range(n):
        f = fila_h + 1 + i
        ws.cell(f, 1, i + 1).border = borde
        ws.cell(f, 2, nombre_al_azar(rnd, i)).border = borde
        ws.cell(f, 3, rnd.choice(["6A", "6B", "7A", "8C", "9B", "10A", "11A"]))
        ws.cell(f, 4, round(rnd.uniform(1, 5), 1)).number_format = "0.0"
        ws.cell(f, 5, round(rnd.uniform(1, 5), 1)).number_format = "0.0"
        ws.cell(f, 6, f"=ROUND(AVERAGE(D{f}:E{f}),1)").number_format = "0.0"
        ws.cell(f, 7, f'=IF(F{f}>=3,"APROBADO","PENDIENTE")')
        if i % 2:
            for c in range(1, 8):
                ws.cell(f, c).fill = PatternFill("solid", fgColor="EEF2FF")
    ult = fila_h + n
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["G"].width = 14
    ws.freeze_panes = f"C{fila_h + 1}"
    ws.auto_filter.ref = f"A{fila_h}:G{ult}"
    ws.conditional_formatting.add(f"F{fila_h + 1}:F{ult}",
                                  CellIsRule(operator="lessThan", formula=["3"], fill=PatternFill("solid", fgColor="FECACA")))
    ws.conditional_formatting.add(f"D{fila_h + 1}:E{ult}", ColorScaleRule(start_type="min", start_color="F8696B",
                                                                          end_type="max", end_color="63BE7B"))
    dv = DataValidation(type="list", formula1='"6A,6B,7A,8C,9B,10A,11A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C{fila_h + 1}:C{ult}")
    ws["B5"].comment = Comment("Comentario de prueba", "Prueba")
    ws["F3"] = "Promedio general:"
    ws["G3"] = f"=AVERAGE(F{fila_h + 1}:F{ult})"

    r = wb.create_sheet("Resumen")
    r["A1"], r["B1"] = "Total estudiantes", f"=COUNTA(Estudiantes!B{fila_h + 1}:B{ult})"
    r["A2"], r["B2"] = "Promedio", f"=AVERAGE(Estudiantes!F{fila_h + 1}:F{ult})"
    r["A3"], r["B3"] = "Aprobados", f'=COUNTIF(Estudiantes!G{fila_h + 1}:G{ult},"APROBADO")'
    r["A4"], r["B4"] = "Promedio (nombre definido)", "=AVERAGE(Promedios)"
    wb.defined_names["Promedios"] = DefinedName("Promedios", attr_text=f"Estudiantes!$F${fila_h + 1}:$F${ult}")
    graf = BarChart()
    graf.title = "Notas (primeros 10)"
    graf.add_data(Reference(ws, min_col=4, max_col=5, min_row=fila_h, max_row=min(ult, fila_h + 10)), titles_from_data=True)
    graf.set_categories(Reference(ws, min_col=1, min_row=fila_h + 1, max_row=min(ult, fila_h + 10)))
    r.add_chart(graf, "D2")
    wb.save(destino)


def libro_tabla(n, destino):
    """Nombres en la primera columna dentro de una tabla de Excel con estilo y fila de totales."""
    rnd = random.Random(1000 + n)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista"
    ws.append(["NOMBRES Y APELLIDOS", "DOCUMENTO", "EDAD"])
    for i in range(n):
        ws.append([nombre_al_azar(rnd, i), 1000000 + i, rnd.randint(10, 18)])
    ws.append(["Total", None, f"=SUBTOTAL(101,C2:C{n + 1})"])
    t = Table(displayName="TablaEstudiantes", ref=f"A1:C{n + 2}", totalsRowCount=1)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(t)
    ws.column_dimensions["A"].width = 40
    wb.save(destino)


def libro_simple(n, destino):
    """Solo una columna con nombres, encabezado en la fila 1."""
    rnd = random.Random(2000 + n)
    wb = Workbook()
    ws = wb.active
    ws.append(["Nombre"])
    for i in range(n):
        ws.append([nombre_al_azar(rnd, i)])
    wb.save(destino)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "pruebas/excels")
    out.mkdir(parents=True, exist_ok=True)
    for n in (100, 500, 1500, 5000):
        libro_completo(n, out / f"completo_{n}.xlsx")
    libro_tabla(300, out / "tabla_300.xlsx")
    libro_simple(50, out / "simple_50.xlsx")
    print("\n".join(sorted(p.name for p in out.glob("*.xlsx"))))
